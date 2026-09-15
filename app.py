import os, calendar, io, secrets, json, hashlib, urllib.request
from datetime import datetime, date, timedelta
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, abort, jsonify, send_file, has_request_context
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import or_, text, event
from sqlalchemy.orm import Session as OrmSession, with_loader_criteria
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

load_dotenv()
app=Flask(__name__)
database_url=os.environ.get('DATABASE_URL','sqlite:///gomusin.db')
if database_url.startswith('postgres://'): database_url=database_url.replace('postgres://','postgresql://',1)
if database_url.startswith('postgresql://'): database_url=database_url.replace('postgresql://','postgresql+psycopg://',1)
app.config.update(
 SECRET_KEY=os.environ.get('SECRET_KEY','CHANGE_THIS_SECRET_KEY'),
 SQLALCHEMY_DATABASE_URI=database_url,
 SQLALCHEMY_TRACK_MODIFICATIONS=False,
 SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE','false').lower()=='true',
 SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
 SQLALCHEMY_ENGINE_OPTIONS={'pool_pre_ping':True},
 MAX_CONTENT_LENGTH=12*1024*1024
)
db=SQLAlchemy(app)

LOGIN_STORIES=[
 ('TRUST IN MOTION','신뢰가 흐르면,','업무는 더 정확해집니다.','고객과의 첫 약속부터 사후관리까지 놓치지 않는 하나의 업무 흐름.',['약속 관리','정확한 기록','안전한 업무']),
 ('CONNECTED GROWTH','연결된 데이터가,','매장의 성장을 만듭니다.','흩어진 고객·판매·재고 정보를 한곳에서 빠르고 분명하게 관리하세요.',['고객 연결','실시간 현황','성장 데이터']),
 ('ONE CLEAR FLOW','복잡한 업무를,','하나의 흐름으로.','오늘 해야 할 일과 중요한 숫자를 가장 먼저 보여드립니다.',['오늘의 업무','빠른 판단','간결한 실행']),
 ('CUSTOMER FIRST','고객을 기억하는,','가장 스마트한 방법.','상담 내용과 약속, 가족 고객의 이력까지 자연스럽게 이어집니다.',['상담 이력','가족 연결','맞춤 응대']),
 ('PRECISION DAILY','매일의 정확함이,','오래가는 신뢰가 됩니다.','작은 누락도 줄이고 모든 업무의 진행 상태를 투명하게 확인하세요.',['누락 방지','진행 추적','책임 관리']),
 ('SECURE BY DESIGN','보이지 않는 곳까지,','안전하게 설계했습니다.','회사와 지점 권한을 서버에서 분리해 필요한 정보만 안전하게 보여줍니다.',['회사별 분리','지점 권한','접근 기록']),
 ('SMART OPERATION','더 적게 찾고,','더 빠르게 처리하세요.','고객 검색부터 개통·재고·정산까지 실무 동선을 짧게 연결합니다.',['통합 검색','자동 입력','빠른 처리']),
 ('PROMISE KEEPER','고객과의 약속을,','끝까지 지킵니다.','변경일·해지일·페이백 일정을 놓치지 않도록 먼저 알려드립니다.',['기한 알림','약속 완료','사후관리']),
 ('VISIBLE RESULT','과정은 투명하게,','결과는 한눈에.','매장별 실적과 정산 흐름을 정확한 숫자로 확인하세요.',['매장 현황','정산 흐름','결과 확인']),
 ('TEAM ALIGNMENT','팀의 모든 업무가,','같은 방향으로 흐릅니다.','담당자와 진행 상황을 공유해 인수인계와 협업을 더 매끄럽게 만듭니다.',['담당자 배정','업무 공유','인수인계']),
 ('DATA WITH PURPOSE','기록을 넘어,','판단에 도움이 되는 데이터.','필요한 순간에 필요한 고객과 업무 정보를 바로 찾을 수 있습니다.',['스마트 검색','이력 연결','업무 판단']),
 ('RELIABLE CONTROL','매장 운영의 중심을,','더 단단하게.','시재·재고·페이백을 기준과 승인 절차에 맞춰 관리합니다.',['시재 관리','재고 추적','승인 절차']),
 ('SEAMLESS SERVICE','상담에서 완료까지,','끊김 없는 고객 경험.','문의와 상담 기록을 이어 받아 누구나 일관된 응대를 할 수 있습니다.',['상담 연속성','응대 기준','고객 경험']),
 ('BRANCH INTELLIGENCE','모든 지점은 연결하고,','권한은 정확히 나눕니다.','관리자는 전체를 보고 직원은 소속 지점 업무에 집중합니다.',['지점 통합','권한 분리','관리자 시야']),
 ('MOMENTUM','오늘의 실행이,','내일의 성장을 앞당깁니다.','우선순위가 분명한 화면으로 중요한 업무부터 빠르게 끝내세요.',['업무 우선순위','집중 실행','성장 리듬']),
 ('CLEAR STANDARD','사람이 바뀌어도,','업무 기준은 흔들리지 않게.','정해진 절차와 기록으로 매장 서비스의 품질을 일정하게 유지합니다.',['표준 업무','품질 유지','안전한 기록']),
 ('TRUSTED INSIGHT','숫자 속에서,','다음 기회를 발견합니다.','판매와 고객 데이터를 연결해 놓치기 쉬운 기회를 보여드립니다.',['판매 분석','고객 기회','실행 제안']),
 ('CALM CONTROL','바쁜 매장일수록,','화면은 더 차분하게.','필요한 정보만 선명하게 정리해 실수와 피로를 줄입니다.',['직관적 화면','실수 예방','업무 집중']),
 ('FUTURE READY','오늘의 매장에서,','내일의 시스템으로.','지점이 늘어나도 같은 기준으로 확장할 수 있는 운영 기반을 만듭니다.',['확장 준비','통합 기준','지속 성장']),
 ('BUILT ON TRUST','신뢰로 연결하고,','데이터로 성장합니다.','고객과의 약속부터 매장 운영까지 하나의 흐름으로 정확하게 관리하세요.',['안전한 계정','실시간 운영','고객 신뢰'])
]

def login_story():
 index=secrets.randbelow(len(LOGIN_STORIES)); kicker,title_a,title_b,body,tags=LOGIN_STORIES[index]
 return {'image':f'images/login/hero-{index+1:02d}.webp','kicker':kicker,'title_a':title_a,'title_b':title_b,'body':body,'tags':tags}

class User(db.Model):
 id=db.Column(db.Integer,primary_key=True); username=db.Column(db.String(50),unique=True,nullable=False,index=True)
 password_hash=db.Column(db.String(255),nullable=False); role=db.Column(db.String(20),nullable=False,default='staff')
 display_name=db.Column(db.String(50)); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id')); active=db.Column(db.Boolean,default=True,nullable=False)
 company_code=db.Column(db.String(50),default='trustflow',nullable=False,index=True); recovery_phone=db.Column(db.String(30))
 can_approve_payback=db.Column(db.Boolean,default=False,nullable=False,index=True)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class AccountRequest(db.Model):
 id=db.Column(db.Integer,primary_key=True); request_type=db.Column(db.String(20),nullable=False)
 company_code=db.Column(db.String(50),nullable=False,index=True); username=db.Column(db.String(50)); display_name=db.Column(db.String(50)); phone=db.Column(db.String(30)); status=db.Column(db.String(20),default='대기',nullable=False); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class PhoneVerification(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 purpose=db.Column(db.String(30),nullable=False,index=True); company_code=db.Column(db.String(50),nullable=False,index=True)
 phone=db.Column(db.String(30),nullable=False,index=True); code_hash=db.Column(db.String(64),nullable=False)
 attempts=db.Column(db.Integer,default=0,nullable=False); verified_at=db.Column(db.DateTime); expires_at=db.Column(db.DateTime,nullable=False,index=True)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True)

class LoginAttempt(db.Model):
 id=db.Column(db.Integer,primary_key=True); company_code=db.Column(db.String(50),nullable=False,index=True); username=db.Column(db.String(50),nullable=False,index=True)
 ip_address=db.Column(db.String(80),nullable=False,index=True); succeeded=db.Column(db.Boolean,default=False,nullable=False); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True)

class Branch(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(100),unique=True,nullable=False,index=True)
 code=db.Column(db.String(30),unique=True); address=db.Column(db.String(255)); phone=db.Column(db.String(30)); manager_name=db.Column(db.String(50))
 company_code=db.Column(db.String(50),default='trustflow',nullable=False,index=True); active=db.Column(db.Boolean,default=True,nullable=False); memo=db.Column(db.Text); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
class Customer(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(100),nullable=False); phone=db.Column(db.String(30),index=True); device=db.Column(db.String(100)); carrier=db.Column(db.String(30)); status=db.Column(db.String(30),default='상담중',nullable=False); memo=db.Column(db.Text)
 company_code=db.Column(db.String(50),default='trustflow',nullable=False,index=True); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'),index=True)
 address_road=db.Column(db.String(255),index=True); address_jibun=db.Column(db.String(255),index=True); address_detail=db.Column(db.String(255)); address_key=db.Column(db.String(255),index=True)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False); updated_at=db.Column(db.DateTime,default=datetime.utcnow,onupdate=datetime.utcnow,nullable=False)
class Booking(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(100),nullable=False); phone=db.Column(db.String(30)); visit_date=db.Column(db.String(50)); device=db.Column(db.String(100)); memo=db.Column(db.Text); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
class Price(db.Model):
 id=db.Column(db.Integer,primary_key=True); device=db.Column(db.String(100),nullable=False); carrier=db.Column(db.String(30),nullable=False); sale_type=db.Column(db.String(30),nullable=False); price=db.Column(db.String(100),nullable=False); updated_at=db.Column(db.DateTime,default=datetime.utcnow,onupdate=datetime.utcnow,nullable=False)
class Partner(db.Model):
 id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(100),unique=True,nullable=False,index=True); category=db.Column(db.String(30)); contact_name=db.Column(db.String(50)); phone=db.Column(db.String(30)); settlement_cycle=db.Column(db.String(30)); default_tax_rate=db.Column(db.Float,default=0.133); active=db.Column(db.Boolean,default=True,nullable=False); memo=db.Column(db.Text); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
class Inventory(db.Model):
 id=db.Column(db.Integer,primary_key=True); serial_number=db.Column(db.String(100),unique=True,nullable=False,index=True); partner_id=db.Column(db.Integer,db.ForeignKey('partner.id')); carrier=db.Column(db.String(30)); manufacturer=db.Column(db.String(50)); model=db.Column(db.String(100),nullable=False); capacity=db.Column(db.String(50)); color=db.Column(db.String(50)); received_date=db.Column(db.Date,default=date.today,nullable=False); purchase_price=db.Column(db.Integer,default=0); storage_location=db.Column(db.String(50)); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id')); status=db.Column(db.String(30),default='보유중',nullable=False,index=True); sale_id=db.Column(db.Integer,index=True); memo=db.Column(db.Text); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
class Sale(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 customer_name=db.Column(db.String(100),nullable=False); customer_phone=db.Column(db.String(30)); customer_birth=db.Column(db.String(20)); customer_gender=db.Column(db.String(10)); opening_date=db.Column(db.Date,default=date.today,nullable=False); carrier=db.Column(db.String(30)); opening_type=db.Column(db.String(30)); status=db.Column(db.String(30),default='개통완료',nullable=False); opening_number=db.Column(db.String(50)); manufacturer=db.Column(db.String(50)); device=db.Column(db.String(100)); color=db.Column(db.String(50)); storage=db.Column(db.String(50)); imei=db.Column(db.String(100)); serial_number=db.Column(db.String(100)); plan=db.Column(db.String(100)); contract_type=db.Column(db.String(50)); installment_months=db.Column(db.String(20)); selection_discount=db.Column(db.String(20)); device_price=db.Column(db.String(50)); official_subsidy=db.Column(db.String(50)); additional_subsidy=db.Column(db.String(50)); seller_subsidy=db.Column(db.String(50)); subsidy=db.Column(db.String(50)); installment_price=db.Column(db.String(50)); cash_price=db.Column(db.String(50)); monthly_installment=db.Column(db.String(50)); monthly_payment=db.Column(db.String(50)); settlement=db.Column(db.String(50)); margin=db.Column(db.String(50)); additional_services=db.Column(db.Text); service_period=db.Column(db.String(50)); gifts=db.Column(db.Text); gift_status=db.Column(db.String(30)); aftercare_status=db.Column(db.String(50)); old_device=db.Column(db.String(100)); old_device_return=db.Column(db.String(20)); trade_in_price=db.Column(db.String(50)); assigned_staff=db.Column(db.String(50)); created_by=db.Column(db.String(50)); memo=db.Column(db.Text)
 partner_id=db.Column(db.Integer,db.ForeignKey('partner.id')); inventory_id=db.Column(db.Integer,db.ForeignKey('inventory.id')); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id')); visit_source=db.Column(db.String(50)); current_plan=db.Column(db.String(100)); next_plan=db.Column(db.String(100)); plan_change_due_date=db.Column(db.Date); rebate=db.Column(db.Integer,default=0); verbal_extra=db.Column(db.Integer,default=0); deduction=db.Column(db.Integer,default=0); extra_support=db.Column(db.Integer,default=0); settlement_amount_v2=db.Column(db.Integer,default=0); tax_rate=db.Column(db.Float,default=0.133); tax_amount=db.Column(db.Integer,default=0); customer_payback=db.Column(db.Integer,default=0); transfer_fee=db.Column(db.Integer,default=0); sim_payment_type=db.Column(db.String(20),default='없음'); sim_fee=db.Column(db.Integer,default=7700); final_margin=db.Column(db.Integer,default=0); internet_carrier=db.Column(db.String(30)); internet_subscriber=db.Column(db.String(100)); internet_install_date=db.Column(db.Date); internet_cancel_due_date=db.Column(db.Date); payback_due_date=db.Column(db.Date); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False); updated_at=db.Column(db.DateTime,default=datetime.utcnow,onupdate=datetime.utcnow,nullable=False)
class SaleAddon(db.Model):
 id=db.Column(db.Integer,primary_key=True); sale_id=db.Column(db.Integer,db.ForeignKey('sale.id',ondelete='CASCADE'),nullable=False,index=True); name=db.Column(db.String(100),nullable=False); retention_rule=db.Column(db.String(30)); cancellation_due_date=db.Column(db.Date); created_at=db.Column(db.DateTime,default=datetime.utcnow)
class CustomerTask(db.Model):
 id=db.Column(db.Integer,primary_key=True); customer_id=db.Column(db.Integer,db.ForeignKey('customer.id')); sale_id=db.Column(db.Integer,db.ForeignKey('sale.id')); task_type=db.Column(db.String(50),nullable=False,index=True); title=db.Column(db.String(150),nullable=False); description=db.Column(db.Text); due_date=db.Column(db.Date,nullable=False,index=True); assigned_staff=db.Column(db.String(50)); status=db.Column(db.String(30),default='처리예정',nullable=False,index=True); auto_created=db.Column(db.Boolean,default=False,nullable=False); completed_at=db.Column(db.DateTime); completed_by=db.Column(db.String(50)); result_memo=db.Column(db.Text); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
class Payback(db.Model):
 id=db.Column(db.Integer,primary_key=True); sale_id=db.Column(db.Integer,db.ForeignKey('sale.id'),nullable=False,index=True); customer_id=db.Column(db.Integer,db.ForeignKey('customer.id')); amount=db.Column(db.Integer,default=0,nullable=False); due_date=db.Column(db.Date,index=True); status=db.Column(db.String(30),default='처리예정',nullable=False,index=True); collection_source=db.Column(db.String(100)); bank=db.Column(db.String(50)); account_number=db.Column(db.String(100)); account_holder=db.Column(db.String(100)); memo=db.Column(db.Text); approval_status=db.Column(db.String(20),default='승인대기',nullable=False,index=True); approved_at=db.Column(db.DateTime); approved_by=db.Column(db.String(50)); rejection_reason=db.Column(db.Text); processed_at=db.Column(db.DateTime); processed_by=db.Column(db.String(50)); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class AuditLog(db.Model):
 id=db.Column(db.Integer,primary_key=True); company_code=db.Column(db.String(50),nullable=False,index=True); branch_id=db.Column(db.Integer,index=True); user_id=db.Column(db.Integer,index=True); username=db.Column(db.String(50),index=True)
 action=db.Column(db.String(50),nullable=False,index=True); target_type=db.Column(db.String(50),index=True); target_id=db.Column(db.String(100)); detail=db.Column(db.Text); ip_address=db.Column(db.String(80)); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True)

class WiredSale(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 sale_date=db.Column(db.Date,nullable=False,index=True)
 customer_name=db.Column(db.String(100),nullable=False,index=True)
 customer_phone=db.Column(db.String(30),index=True)
 subscriber_name=db.Column(db.String(100))
 branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'),index=True)
 assigned_staff=db.Column(db.String(50),index=True)
 carrier=db.Column(db.String(30),index=True)
 business_type=db.Column(db.String(30),default='유선판매',nullable=False,index=True)
 product_type=db.Column(db.String(50),index=True)
 internet_plan=db.Column(db.String(120))
 internet_speed=db.Column(db.String(20))
 tv_plan=db.Column(db.String(100))
 install_due_date=db.Column(db.Date,index=True)
 install_date=db.Column(db.Date,index=True)
 status=db.Column(db.String(30),default='접수',nullable=False,index=True)
 rebate=db.Column(db.Integer,default=0)
 verbal_extra=db.Column(db.Integer,default=0)
 gift_certificate=db.Column(db.Integer,default=0)
 gift_cost=db.Column(db.Integer,default=0)
 deduction=db.Column(db.Integer,default=0)
 settlement_amount=db.Column(db.Integer,default=0)
 tax_rate=db.Column(db.Float,default=0.133)
 tax_amount=db.Column(db.Integer,default=0)
 payback=db.Column(db.Integer,default=0)
 gift_return=db.Column(db.Integer,default=0)
 final_margin=db.Column(db.Integer,default=0)
 memo=db.Column(db.Text)
 created_by=db.Column(db.String(50)); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False); updated_at=db.Column(db.DateTime,default=datetime.utcnow,onupdate=datetime.utcnow)

class CashLedger(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 ledger_date=db.Column(db.Date,nullable=False,index=True); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'),nullable=False,index=True)
 direction=db.Column(db.String(10),nullable=False,index=True)  # 입금 / 출금
 category=db.Column(db.String(50),nullable=False,index=True); amount=db.Column(db.Integer,default=0,nullable=False)
 payment_method=db.Column(db.String(30),default='현금',nullable=False); reference_type=db.Column(db.String(30)); reference_id=db.Column(db.Integer)
 counterparty=db.Column(db.String(100)); memo=db.Column(db.Text); created_by=db.Column(db.String(50)); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True)

class CardTerminal(db.Model):
 id=db.Column(db.Integer,primary_key=True); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'),nullable=False,index=True)
 provider=db.Column(db.String(50)); merchant_number=db.Column(db.String(100),index=True); terminal_number=db.Column(db.String(100),unique=True,nullable=False,index=True)
 api_token=db.Column(db.String(120),nullable=False,unique=True,index=True); active=db.Column(db.Boolean,default=True,nullable=False); memo=db.Column(db.Text)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class CardTransaction(db.Model):
 id=db.Column(db.Integer,primary_key=True); terminal_id=db.Column(db.Integer,db.ForeignKey('card_terminal.id'),nullable=False,index=True); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'),nullable=False,index=True)
 approval_number=db.Column(db.String(100),nullable=False,index=True); paid_at=db.Column(db.DateTime,nullable=False,index=True); amount=db.Column(db.Integer,default=0,nullable=False)
 card_company=db.Column(db.String(50)); installment=db.Column(db.String(20)); receipt_number=db.Column(db.String(100)); status=db.Column(db.String(20),default='승인',nullable=False,index=True)
 raw_data=db.Column(db.Text); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
 __table_args__=(db.UniqueConstraint('terminal_id','approval_number',name='uq_terminal_approval'),)

class ContactLog(db.Model):
 id=db.Column(db.Integer,primary_key=True); customer_id=db.Column(db.Integer,db.ForeignKey('customer.id'),nullable=False,index=True); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'),nullable=False,index=True)
 contacted_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True); staff_name=db.Column(db.String(50),nullable=False,index=True)
 channel=db.Column(db.String(30),default='전화'); outcome=db.Column(db.String(30),default='상담완료',index=True); note=db.Column(db.Text,nullable=False); next_contact_date=db.Column(db.Date,index=True)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class LegalCase(db.Model):
 id=db.Column(db.Integer,primary_key=True); customer_id=db.Column(db.Integer,db.ForeignKey('customer.id'),nullable=False,index=True); branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'),nullable=False,index=True)
 case_type=db.Column(db.String(40),default='환수',nullable=False,index=True); claim_amount=db.Column(db.Integer,default=0,nullable=False); incident_date=db.Column(db.Date)
 reason=db.Column(db.Text); evidence=db.Column(db.Text); debtor_address=db.Column(db.String(300)); demand_due_date=db.Column(db.Date,index=True)
 status=db.Column(db.String(30),default='자료수집',nullable=False,index=True); assigned_staff=db.Column(db.String(50)); memo=db.Column(db.Text)
 created_by=db.Column(db.String(50)); created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True); updated_at=db.Column(db.DateTime,default=datetime.utcnow,onupdate=datetime.utcnow)


class PlanMaster(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 carrier=db.Column(db.String(20),nullable=False,index=True)
 name=db.Column(db.String(120),nullable=False,index=True)
 active=db.Column(db.Boolean,default=True,nullable=False,index=True)
 sort_order=db.Column(db.Integer,default=0)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class WiredProductMaster(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 carrier=db.Column(db.String(20),nullable=False,index=True)
 category=db.Column(db.String(20),nullable=False,index=True)  # internet / tv
 name=db.Column(db.String(120),nullable=False,index=True)
 speed=db.Column(db.String(20))
 active=db.Column(db.Boolean,default=True,nullable=False,index=True)
 sort_order=db.Column(db.Integer,default=0)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class DeviceMaster(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 manufacturer=db.Column(db.String(30),nullable=False,index=True)
 model=db.Column(db.String(120),nullable=False,index=True)
 capacities=db.Column(db.String(200))
 colors=db.Column(db.String(500))
 active=db.Column(db.Boolean,default=True,nullable=False,index=True)
 sort_order=db.Column(db.Integer,default=0)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class SaleDocument(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 sale_id=db.Column(db.Integer,db.ForeignKey('sale.id',ondelete='CASCADE'),nullable=False,index=True)
 branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'),index=True)
 doc_type=db.Column(db.String(50),default='기타서류')
 original_name=db.Column(db.String(255),nullable=False)
 content_type=db.Column(db.String(100),nullable=False)
 file_size=db.Column(db.Integer,default=0)
 file_data=db.Column(db.LargeBinary,nullable=False)
 uploaded_by=db.Column(db.String(50))
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True)

class InventoryMovement(db.Model):
 id=db.Column(db.Integer,primary_key=True)
 inventory_id=db.Column(db.Integer,db.ForeignKey('inventory.id'),nullable=False,index=True)
 action=db.Column(db.String(30),nullable=False,index=True)
 from_branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'))
 to_branch_id=db.Column(db.Integer,db.ForeignKey('branch.id'))
 from_status=db.Column(db.String(30)); to_status=db.Column(db.String(30))
 processed_by=db.Column(db.String(50)); memo=db.Column(db.Text)
 created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True)

def money(v):
 try:return int(float(str(v or 0).replace(',','').replace('원','').strip() or 0))
 except:return 0

def normalize_phone(v):
 digits=''.join(ch for ch in str(v or '') if ch.isdigit())
 return digits[:11]

def _verification_hash(code):
 return hashlib.sha256(f"{app.config['SECRET_KEY']}:{code}".encode()).hexdigest()

def _send_sms(phone,message):
 """Send through a provider-neutral HTTPS webhook configured in Render."""
 if app.config.get('TESTING'):return True
 endpoint=os.environ.get('SMS_WEBHOOK_URL','').strip(); token=os.environ.get('SMS_WEBHOOK_TOKEN','').strip()
 if not endpoint:return False
 payload=json.dumps({'to':phone,'message':message,'sender':os.environ.get('SMS_SENDER','TrustFlow')},ensure_ascii=False).encode()
 headers={'Content-Type':'application/json'}
 if token:headers['Authorization']=f'Bearer {token}'
 try:
  with urllib.request.urlopen(urllib.request.Request(endpoint,data=payload,headers=headers,method='POST'),timeout=8) as response:
   return 200<=response.status<300
 except Exception:return False

def issue_phone_code(purpose,company,phone):
 now=datetime.utcnow(); recent=PhoneVerification.query.filter_by(purpose=purpose,company_code=company,phone=phone).filter(PhoneVerification.created_at>now-timedelta(minutes=1)).first()
 if recent:return False,'인증번호는 1분 후 다시 요청할 수 있습니다.'
 code=os.environ.get('SMS_TEST_CODE','123456') if app.config.get('TESTING') else f'{secrets.randbelow(1000000):06d}'
 item=PhoneVerification(purpose=purpose,company_code=company,phone=phone,code_hash=_verification_hash(code),expires_at=now+timedelta(minutes=5))
 db.session.add(item);db.session.commit()
 if not _send_sms(phone,f'[TrustFlow] 인증번호는 {code}입니다. 5분 안에 입력해주세요.'):
  db.session.delete(item);db.session.commit();return False,'문자 인증 서비스 연결이 아직 완료되지 않았습니다. 관리자에게 문의해주세요.'
 return True,'인증번호를 문자로 보냈습니다. 5분 안에 입력해주세요.'

def verify_phone_code(purpose,company,phone,code):
 item=PhoneVerification.query.filter_by(purpose=purpose,company_code=company,phone=phone).order_by(PhoneVerification.id.desc()).first(); now=datetime.utcnow()
 if not item or item.verified_at or item.expires_at<now:return False,'인증번호가 만료됐습니다. 다시 받아주세요.'
 if item.attempts>=5:return False,'입력 횟수를 초과했습니다. 새 인증번호를 받아주세요.'
 item.attempts+=1
 if not secrets.compare_digest(item.code_hash,_verification_hash((code or '').strip())):
  db.session.commit();return False,'인증번호가 올바르지 않습니다.'
 item.verified_at=now;db.session.commit();return True,''

def parse_date(v):
 try:return datetime.strptime((v or '').strip(),'%Y-%m-%d').date() if v else None
 except:return None

def add_months(d,n):
 m=d.month-1+n; y=d.year+m//12; m=m%12+1; return date(y,m,min(d.day,calendar.monthrange(y,m)[1]))
def due_from_rule(d,rule):
 rule=(rule or '').upper().strip()
 if rule.startswith('D+'):
  try:return d+timedelta(days=int(rule[2:]))
  except:return None
 if rule.startswith('M+'):
  try:return add_months(d,int(rule[2:]))
  except:return None
 return parse_date(rule)
def calc_settlement(rebate,verbal,deduct,support,payback,opening_type='',sim_payment_type='없음',tax_rate=.133):
 base=rebate+verbal-deduct-support
 transfer_fee=800 if opening_type=='번호이동' else 0
 settlement=base-transfer_fee
 tax=round(settlement*tax_rate)
 sim_cost=7700 if sim_payment_type=='선납' else 0
 return settlement,tax,settlement-tax-payback-sim_cost,transfer_fee

def calc_wired_settlement(rebate,verbal,gift_certificate,gift_cost,deduction,payback,gift_return,tax_rate=.133):
 # 유선 정산: 리베이트 + 구두추가 + 정산상품권 - 고객지급상품권 - 차감금액
 settlement=rebate+verbal+gift_certificate-gift_cost-deduction
 tax=round(settlement*tax_rate)
 # 상품권 반납은 회수된 금액이므로 최종마진에 더합니다.
 final_margin=settlement-tax-payback+gift_return
 return settlement,tax,final_margin


def current_branch_id():
 try:return int(session.get('branch_id')) if session.get('branch_id') not in [None,''] else None
 except:return None

def current_company():
 return (session.get('company_code') or 'trustflow').strip().lower()

@event.listens_for(OrmSession,'do_orm_execute')
def tenant_read_filter(execute_state):
 if execute_state.execution_options.get('skip_tenant') or not execute_state.is_select or not has_request_context() or not session.get('user_id'):return
 company=current_company()
 execute_state.statement=execute_state.statement.options(
  with_loader_criteria(Branch,lambda row:row.company_code==company,include_aliases=True),
  with_loader_criteria(Customer,lambda row:row.company_code==company,include_aliases=True)
 )

@event.listens_for(OrmSession,'before_flush')
def tenant_write_defaults(db_session,flush_context,instances):
 if not has_request_context() or not session.get('user_id'):return
 company=current_company()
 for obj in db_session.new:
  if isinstance(obj,(Branch,Customer)) and not obj.company_code:obj.company_code=company

def is_admin():
 return session.get('role')=='admin'

def enforce_user_company(user):
 if not user or user.company_code!=(session.get('company_code') or 'trustflow'): abort(403)

def enforce_branch(branch_id):
 try:bid=int(branch_id or 0)
 except:abort(403)
 branch=db.session.get(Branch,bid)
 if not branch or branch.company_code!=current_company():abort(403)
 if not is_admin() and (not current_branch_id() or bid!=current_branch_id()): abort(403)

def apply_branch_scope(query, model):
 if is_admin():
  query=query.filter(model.branch_id.in_(db.session.query(Branch.id)))
 else:
  bid=current_branch_id()
  if not bid:return query.filter(db.text('1=0'))
  query=query.filter(model.branch_id==bid)
 return query

def customer_query_scoped():
 q=Customer.query
 if is_admin():return q
 bid=current_branch_id()
 if not bid:return q.filter(Customer.id==-1)
 branch_phones=db.session.query(Sale.customer_phone).filter(Sale.branch_id==bid,Sale.customer_phone.isnot(None))
 return q.filter(or_(Customer.branch_id==bid,Customer.phone.in_(branch_phones)))

def customer_allowed(customer):
 if not customer:return False
 if is_admin():return True
 bid=current_branch_id()
 return bool(bid and (customer.branch_id==bid or (customer.phone and Sale.query.filter_by(customer_phone=customer.phone,branch_id=bid).first())))

def sale_allowed(sale):
 return bool(sale and (is_admin() or (current_branch_id() and sale.branch_id==current_branch_id())))

def task_query_scoped():
 q=CustomerTask.query.outerjoin(Sale,CustomerTask.sale_id==Sale.id)
 if not is_admin():
  bid=current_branch_id()
  q=q.filter(Sale.branch_id==bid) if bid else q.filter(CustomerTask.id==-1)
 return q

def payback_query_scoped():
 q=Payback.query.join(Sale,Payback.sale_id==Sale.id)
 if not is_admin():
  bid=current_branch_id()
  q=q.filter(Sale.branch_id==bid) if bid else q.filter(Payback.id==-1)
 return q

def can_approve_payback():
 if is_admin():return True
 user_id=session.get('user_id')
 if not user_id:return False
 try:
  user=db.session.get(User,user_id)
  return bool(user and user.can_approve_payback)
 except:return False

def notification_summary():
 if not session.get('user_id'):return {'total':0,'overdue_tasks':0,'today_tasks':0,'pending_approvals':0,'due_paybacks':0,'legal_deadlines':0}
 today=date.today(); open_states=['처리예정','연락안됨','연기']
 overdue_tasks=task_query_scoped().filter(CustomerTask.due_date<today,CustomerTask.status.in_(open_states)).count()
 today_tasks=task_query_scoped().filter(CustomerTask.due_date==today,CustomerTask.status.in_(open_states)).count()
 pq=payback_query_scoped()
 pending_approvals=pq.filter(Payback.status!='완료',Payback.approval_status=='승인대기').count() if can_approve_payback() else 0
 due_paybacks=pq.filter(Payback.status!='완료',Payback.due_date<=today).count()
 lq=apply_branch_scope(LegalCase.query,LegalCase)
 legal_deadlines=lq.filter(LegalCase.status.notin_(['완료','종결']),LegalCase.demand_due_date.isnot(None),LegalCase.demand_due_date<=today+timedelta(days=3)).count()
 return {'total':overdue_tasks+today_tasks+pending_approvals+due_paybacks+legal_deadlines,'overdue_tasks':overdue_tasks,'today_tasks':today_tasks,'pending_approvals':pending_approvals,'due_paybacks':due_paybacks,'legal_deadlines':legal_deadlines}

def audit(action,target_type='',target_id='',detail='',branch_id=None,commit=False):
 try:
  db.session.add(AuditLog(company_code=session.get('company_code') or 'trustflow',branch_id=branch_id or current_branch_id(),user_id=session.get('user_id'),username=session.get('display_name') or session.get('username') or 'system',action=action,target_type=target_type,target_id=str(target_id or ''),detail=str(detail or '')[:2000],ip_address=(request.headers.get('X-Forwarded-For','').split(',')[0].strip() or request.remote_addr)))
  if commit:db.session.commit()
 except:db.session.rollback()

def login_required(fn):
 @wraps(fn)
 def wrapped(*a,**kw):
  if not session.get('user_id'):return redirect(url_for('login'))
  return fn(*a,**kw)
 return wrapped

def admin_required(fn):
 @wraps(fn)
 def wrapped(*a,**kw):
  if session.get('role')!='admin':abort(403)
  return fn(*a,**kw)
 return wrapped

def _add_columns(table, cols):
 inspector=db.inspect(db.engine)
 if table not in inspector.get_table_names(): return
 existing={c['name'] for c in inspector.get_columns(table)}
 for name,sqltype in cols.items():
  if name not in existing:
   try: db.session.execute(text(f'ALTER TABLE "{table}" ADD COLUMN "{name}" {sqltype}')); db.session.commit()
   except Exception: db.session.rollback()

def upgrade_existing_sale():
 _add_columns('sale',{
  'partner_id':'INTEGER','inventory_id':'INTEGER','branch_id':'INTEGER','visit_source':'VARCHAR(50)',
  'current_plan':'VARCHAR(100)','next_plan':'VARCHAR(100)','plan_change_due_date':'DATE',
  'rebate':'INTEGER DEFAULT 0','verbal_extra':'INTEGER DEFAULT 0','deduction':'INTEGER DEFAULT 0',
  'extra_support':'INTEGER DEFAULT 0','settlement_amount_v2':'INTEGER DEFAULT 0','tax_rate':'FLOAT DEFAULT 0.133',
  'tax_amount':'INTEGER DEFAULT 0','customer_payback':'INTEGER DEFAULT 0','transfer_fee':'INTEGER DEFAULT 0',
  'sim_payment_type':"VARCHAR(20) DEFAULT '없음'",'sim_fee':'INTEGER DEFAULT 7700','final_margin':'INTEGER DEFAULT 0',
  'internet_carrier':'VARCHAR(30)','internet_subscriber':'VARCHAR(100)','internet_install_date':'DATE',
  'internet_cancel_due_date':'DATE','payback_due_date':'DATE','created_at':'TIMESTAMP','updated_at':'TIMESTAMP'
 })
 _add_columns('user',{'display_name':'VARCHAR(50)','branch_id':'INTEGER','active':'BOOLEAN DEFAULT TRUE'})
 _add_columns('inventory',{'branch_id':'INTEGER'})
 _add_columns('wired_sale',{
  'internet_plan':'VARCHAR(120)','internet_speed':'VARCHAR(20)','tv_plan':'VARCHAR(100)',
  'rebate':'INTEGER DEFAULT 0','verbal_extra':'INTEGER DEFAULT 0',
  'gift_certificate':'INTEGER DEFAULT 0','gift_cost':'INTEGER DEFAULT 0','deduction':'INTEGER DEFAULT 0',
  'tax_rate':'FLOAT DEFAULT 0.133','tax_amount':'INTEGER DEFAULT 0',
  'payback':'INTEGER DEFAULT 0','gift_return':'INTEGER DEFAULT 0',
  'final_margin':'INTEGER DEFAULT 0'
 })

def seed_branches():
 for idx,name in enumerate(['1호점','2호점','3호점'],1):
  if not Branch.query.execution_options(skip_tenant=True).filter_by(name=name).first(): db.session.add(Branch(name=name,code=f'B{idx:02d}',company_code='trustflow'))
 db.session.commit()


def seed_masters():
 # 통신사 요금제는 관리자 기준정보 화면에서 언제든 추가/중지 가능
 plan_seed={
  'SK':['베스트 129','베스트 119','베스트 109','베스트 99','베스트 89',
        '라이트 79','라이트 69','라이트 59','라이트 49','라이트 43','라이트 39','T플랜 세이브','ZEM플랜 스마트'],
  'KT':['초이스 프리미엄','초이스 스페셜','초이스 베이직','베이직21GB(이월)','베이직','슬림','5G 주니어'],
  'LG':['플러스플랜130','플러스플랜115','플러스플랜105','플러스플랜95','데이터플랜MAX',
        '5G 프리미어 레귤러','5G 프리미어 에센셜']
 }
 for carrier,names in plan_seed.items():
  for idx,name in enumerate(names):
   if not PlanMaster.query.filter_by(carrier=carrier,name=name).first():
    db.session.add(PlanMaster(carrier=carrier,name=name,active=True,sort_order=idx))

 devices=[
  ('삼성','갤럭시 S26','256GB,512GB','코발트 바이올렛,화이트,블랙,실버'),
  ('삼성','갤럭시 S26+','256GB,512GB','코발트 바이올렛,화이트,블랙,실버'),
  ('삼성','갤럭시 S26 울트라','256GB,512GB,1TB','블랙,화이트,실버,블루'),
  ('삼성','갤럭시 Z 폴드8','256GB,512GB,1TB','블랙,실버,블루'),
  ('삼성','갤럭시 Z 폴드8 울트라','512GB,1TB','블랙,실버'),
  ('삼성','갤럭시 Z 플립8','256GB,512GB','블랙,화이트,블루,핑크'),
  ('삼성','갤럭시 S25','256GB,512GB','네이비,실버 쉐도우,아이스블루,민트'),
  ('삼성','갤럭시 S25+','256GB,512GB','네이비,실버 쉐도우,아이스블루,민트'),
  ('삼성','갤럭시 S25 울트라','256GB,512GB,1TB','티타늄 블랙,티타늄 그레이,티타늄 실버블루'),
  ('삼성','갤럭시 A56','128GB,256GB','블랙,그레이,올리브'),
  ('애플','iPhone 17 Pro Max','256GB,512GB,1TB,2TB','실버,코스믹 오렌지,딥 블루'),
  ('애플','iPhone 17 Pro','256GB,512GB,1TB','실버,코스믹 오렌지,딥 블루'),
  ('애플','iPhone Air','256GB,512GB,1TB','스페이스 블랙,클라우드 화이트,라이트 골드,스카이 블루'),
  ('애플','iPhone 17','256GB,512GB','블랙,화이트,미스트 블루,세이지,라벤더'),
  ('애플','iPhone 17e','128GB,256GB,512GB','블랙,화이트'),
  ('애플','iPhone 16','128GB,256GB,512GB','블랙,화이트,핑크,틸,울트라마린'),
  ('애플','iPhone 16e','128GB,256GB,512GB','블랙,화이트')
 ]
 wired_seed={
  'SK':{
   'internet':[
    ('광랜 (100M)','100M'),
    ('기가라이트 (500M)','500M'),
    ('기가 (1G)','1G')
   ],
   'tv':[
    ('B tv 이코노미',None),
    ('B tv 스탠다드',None),
    ('B tv All',None),
    ('B tv All+지상파',None),
    ('B tv All+캐치온',None),
    ('B tv All+캐치온+지상파',None),
    ('B tv 미니',None)
   ]
  },
  'KT':{
   'internet':[
    ('인터넷 슬림 (100M)','100M'),
    ('인터넷 베이직 (500M)','500M'),
    ('인터넷 에센스 (1G)','1G'),
    ('요고 인터넷 슬림 와이파이 (100M)','100M'),
    ('요고 인터넷 베이직 와이파이 (500M)','500M'),
    ('요고 인터넷 에센스 와이파이 (1G)','1G')
   ],
   'tv':[
    ('지니 TV 베이직',None),
    ('지니 TV 라이트',None),
    ('지니 TV 에센스',None),
    ('지니 TV 모든G',None),
    ('지니 TV 디즈니+ 모든G',None),
    ('지니 TV VOD초이스',None),
    ('지니 TV 슈퍼팩 초이스',None),
    ('지니 TV 넷플릭스 초이스HD',None)
   ]
  },
  'LG':{
   'internet':[
    ('너겟 라이트 100M','100M'),
    ('너겟 라이트 500M','500M'),
    ('너겟 라이트 1G','1G'),
    ('너겟 100M','100M'),
    ('너겟 500M','500M')
   ],
   'tv':[
    ('실속형',None),
    ('기본형',None),
    ('고급형',None),
    ('프리미엄',None),
    ('기본형 방송패스',None),
    ('프리미엄 방송패스',None),
    ('프리미엄 VOD',None),
    ('프리미엄 유플레이',None)
   ]
  }
 }
 for carrier,cats in wired_seed.items():
  for category,rows in cats.items():
   for idx,(name,speed) in enumerate(rows):
    if not WiredProductMaster.query.filter_by(carrier=carrier,category=category,name=name).first():
     db.session.add(WiredProductMaster(carrier=carrier,category=category,name=name,speed=speed,active=True,sort_order=idx))

 for idx,(maker,model,caps,colors) in enumerate(devices):
  if not DeviceMaster.query.filter_by(manufacturer=maker,model=model).first():
   db.session.add(DeviceMaster(manufacturer=maker,model=model,capacities=caps,colors=colors,active=True,sort_order=idx))
 db.session.commit()

def prepare_database():
 db.create_all(); _add_columns('user',{'company_code':"VARCHAR(50) DEFAULT 'trustflow'",'recovery_phone':'VARCHAR(30)','can_approve_payback':'BOOLEAN DEFAULT FALSE'}); _add_columns('branch',{'company_code':"VARCHAR(50) DEFAULT 'trustflow'"}); _add_columns('customer',{'company_code':"VARCHAR(50) DEFAULT 'trustflow'",'branch_id':'INTEGER','address_road':'VARCHAR(255)','address_jibun':'VARCHAR(255)','address_detail':'VARCHAR(255)','address_key':'VARCHAR(255)'}); _add_columns('payback',{'approval_status':"VARCHAR(20) DEFAULT '승인대기'",'approved_at':'TIMESTAMP','approved_by':'VARCHAR(50)','rejection_reason':'TEXT'}); _add_columns('wired_sale',{'business_type':"VARCHAR(30) DEFAULT '유선판매'"}); upgrade_existing_sale(); seed_branches(); seed_masters()

def sync_admin():
 prepare_database(); u=os.environ.get('ADMIN_USERNAME','').strip(); p=os.environ.get('ADMIN_PASSWORD',''); company_code=os.environ.get('COMPANY_LOGIN_ID','trustflow').strip().lower() or 'trustflow'
 if not u or not p:return
 user=User.query.filter_by(username=u).first()
 if not user:db.session.add(User(username=u,password_hash=generate_password_hash(p),role='admin',company_code=company_code));db.session.commit();return
 changed=False
 if not check_password_hash(user.password_hash,p):user.password_hash=generate_password_hash(p);changed=True
 if user.role!='admin':user.role='admin';changed=True
 if not user.company_code:user.company_code=company_code;changed=True
 if changed:db.session.commit()

@app.context_processor
def helpers():return dict(current_user=session.get('display_name') or session.get('username'),current_role=session.get('role'),current_company=session.get('company_code'),current_branch_id=current_branch_id(),can_approve_payback=can_approve_payback(),notification_summary=notification_summary(),moneyfmt=lambda v:f'{money(v):,}')

@app.errorhandler(403)
def forbidden_error(error):
 return render_template('error.html',code=403,title='접근 권한이 없습니다',message='현재 계정 또는 소속 회사에서 사용할 수 없는 메뉴입니다.'),403

@app.errorhandler(404)
def not_found_error(error):
 return render_template('error.html',code=404,title='페이지를 찾을 수 없습니다',message='주소가 변경됐거나 존재하지 않는 화면입니다.'),404

@app.errorhandler(500)
def server_error(error):
 db.session.rollback()
 return render_template('error.html',code=500,title='잠시 처리할 수 없습니다',message='입력한 내용은 다시 확인할 수 있도록 안전하게 처리하고 있습니다.'),500

@app.route('/health')
def health():
 try:db.session.execute(text('SELECT 1'));return {'status':'ok','database':'connected'}
 except Exception as e:return {'status':'error','message':str(e)},500
@app.route('/login',methods=['GET','POST'])
def login():
 try:sync_admin()
 except Exception as e:return f'DB 연결 오류: {e}',500
 if request.method=='POST':
  company_code=request.form.get('company_code','').strip().lower()
  username=request.form.get('username','').strip(); ip=(request.headers.get('X-Forwarded-For','').split(',')[0].strip() or request.remote_addr or 'unknown'); since=datetime.utcnow()-timedelta(minutes=15)
  failures=LoginAttempt.query.filter_by(company_code=company_code,username=username,ip_address=ip,succeeded=False).filter(LoginAttempt.created_at>=since).count()
  if failures>=5:
   flash('로그인 시도가 많습니다. 15분 후 다시 시도하거나 비밀번호를 재설정해주세요.','error');return render_template('login.html',story=login_story()),429
  user=User.query.filter_by(username=username,company_code=company_code).first()
  if user and user.active is False:
   flash('비활성화된 직원 계정입니다. 관리자에게 문의해주세요.','error'); return render_template('login.html',story=login_story())
  if user and check_password_hash(user.password_hash,request.form.get('password','')):
   db.session.add(LoginAttempt(company_code=company_code,username=username,ip_address=ip,succeeded=True));db.session.commit();session.clear();session.update(user_id=user.id,username=user.username,display_name=user.display_name or user.username,role=user.role,branch_id=user.branch_id,company_code=user.company_code);return redirect(url_for('dashboard'))
  db.session.add(LoginAttempt(company_code=company_code,username=username,ip_address=ip,succeeded=False));db.session.commit()
  flash('아이디 또는 비밀번호가 올바르지 않습니다.','error')
 return render_template('login.html',story=login_story())

@app.route('/signup',methods=['GET','POST'])
def signup():
 prepare_database()
 if request.method=='POST':
  company=request.form.get('company_code','').strip().lower(); username=request.form.get('username','').strip(); name=request.form.get('display_name','').strip(); phone=normalize_phone(request.form.get('phone','')); password=request.form.get('password','')
  if not all([company,username,name,phone,password]): flash('모든 항목을 입력해주세요.','error')
  elif not User.query.filter_by(company_code=company).first(): flash('등록되지 않은 회사 전체아이디입니다.','error')
  elif User.query.filter_by(username=username).first(): flash('이미 사용 중인 개인아이디입니다.','error')
  else:
   db.session.add(User(username=username,password_hash=generate_password_hash(password),role='staff',display_name=name,company_code=company,recovery_phone=phone,active=False)); db.session.commit(); flash('가입 신청이 완료됐습니다. 회사 관리자의 승인을 기다려주세요.','success'); return redirect(url_for('login'))
 return render_template('signup.html')

@app.route('/find-id',methods=['GET','POST'])
def find_id():
 prepare_database(); found=None; verification_sent=False
 if request.method=='POST':
  company=request.form.get('company_code','').strip().lower(); name=request.form.get('display_name','').strip(); phone=normalize_phone(request.form.get('phone','')); action=request.form.get('action','send')
  user=User.query.filter_by(company_code=company,display_name=name,recovery_phone=phone).first()
  if action=='send':
   if user:
    ok,message=issue_phone_code('find_id',company,phone); flash(message,'success' if ok else 'error'); verification_sent=ok
   else:flash('입력한 정보와 일치하는 계정을 찾지 못했습니다.','error')
  elif action=='verify' and user:
   ok,message=verify_phone_code('find_id',company,phone,request.form.get('code'))
   if ok:found=user.username
   else:flash(message,'error');verification_sent=True
  else:flash('입력한 정보와 일치하는 계정을 찾지 못했습니다.','error')
 return render_template('find_id.html',found=found,verification_sent=verification_sent,form=request.form)

@app.route('/password-help',methods=['GET','POST'])
def password_help():
 prepare_database(); verification_sent=False; reset_done=False
 if request.method=='POST':
  company=request.form.get('company_code','').strip().lower(); username=request.form.get('username','').strip(); name=request.form.get('display_name','').strip(); phone=normalize_phone(request.form.get('phone','')); action=request.form.get('action','send')
  user=User.query.filter_by(company_code=company,username=username,display_name=name,recovery_phone=phone).first()
  if action=='send':
   if user:
    ok,message=issue_phone_code('password_reset',company,phone);flash(message,'success' if ok else 'error');verification_sent=ok
   else:flash('입력한 정보와 일치하는 계정을 찾지 못했습니다.','error')
  elif action=='reset' and user:
   ok,message=verify_phone_code('password_reset',company,phone,request.form.get('code')); password=request.form.get('new_password','')
   if not ok:flash(message,'error');verification_sent=True
   elif len(password)<8:flash('새 비밀번호는 8자 이상 입력해주세요.','error');verification_sent=True
   else:
    user.password_hash=generate_password_hash(password);db.session.add(AccountRequest(request_type='비밀번호완료',company_code=company,username=username,display_name=name,phone=phone,status='완료'));db.session.commit();reset_done=True
  else:flash('입력한 정보와 일치하는 계정을 찾지 못했습니다.','error')
 return render_template('password_help.html',verification_sent=verification_sent,reset_done=reset_done,form=request.form)

@app.route('/logout')
def logout():session.clear();return redirect(url_for('login'))

def scoped_branch_from_request():
 if not is_admin(): return current_branch_id()
 try:return int(request.values.get('branch_id')) if request.values.get('branch_id') else None
 except:return None

@app.route('/cash-ledger',methods=['GET','POST'])
@login_required
def cash_ledger():
 prepare_database(); branch_id=scoped_branch_from_request(); month=request.values.get('month') or date.today().strftime('%Y-%m')
 try:y,m=map(int,month.split('-')); start=date(y,m,1); end=add_months(start,1)
 except:y,m=date.today().year,date.today().month; start=date(y,m,1); end=add_months(start,1); month=start.strftime('%Y-%m')
 if request.method=='POST':
  bid=current_branch_id() if not is_admin() else (request.form.get('branch_id') or None)
  if not bid: flash('지점을 선택해주세요.','error'); return redirect(url_for('cash_ledger',month=month))
  direction=request.form.get('direction','입금'); amount=abs(money(request.form.get('amount')))
  if amount<=0: flash('금액을 입력해주세요.','error'); return redirect(url_for('cash_ledger',month=month,branch_id=bid))
  db.session.add(CashLedger(ledger_date=parse_date(request.form.get('ledger_date')) or date.today(),branch_id=int(bid),direction=direction,category=request.form.get('category','기타'),amount=amount,payment_method=request.form.get('payment_method','현금'),counterparty=request.form.get('counterparty'),memo=request.form.get('memo'),created_by=session.get('display_name') or session.get('username')))
  db.session.commit(); flash('시재 내역이 등록되었습니다.','success'); return redirect(url_for('cash_ledger',month=month,branch_id=bid))
 q=CashLedger.query.filter(CashLedger.ledger_date>=start,CashLedger.ledger_date<end)
 if branch_id:q=q.filter_by(branch_id=branch_id)
 elif not is_admin():q=q.filter(CashLedger.id==-1)
 items=q.order_by(CashLedger.ledger_date.desc(),CashLedger.id.desc()).all(); branch_map={b.id:b for b in Branch.query.all()}
 cash_in=sum(x.amount for x in items if x.direction=='입금' and x.payment_method=='현금'); cash_out=sum(x.amount for x in items if x.direction=='출금' and x.payment_method=='현금'); card_total=sum(x.amount for x in items if x.direction=='입금' and x.payment_method=='카드')-sum(x.amount for x in items if x.direction=='출금' and x.payment_method=='카드')
 return render_template('cash_ledger.html',items=items,month=month,branch_id=branch_id,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all(),branch_map=branch_map,cash_in=cash_in,cash_out=cash_out,cash_balance=cash_in-cash_out,card_total=card_total,today=date.today().isoformat())

@app.get('/cash-ledger/export')
@login_required
def cash_ledger_export():
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 month=request.args.get('month') or date.today().strftime('%Y-%m'); branch_id=scoped_branch_from_request()
 try:y,m=map(int,month.split('-')); start=date(y,m,1); end=add_months(start,1)
 except:abort(400)
 q=CashLedger.query.filter(CashLedger.ledger_date>=start,CashLedger.ledger_date<end)
 if branch_id:q=q.filter_by(branch_id=branch_id)
 elif not is_admin():q=q.filter(CashLedger.id==-1)
 items=q.order_by(CashLedger.ledger_date,CashLedger.id).all(); branches={b.id:b.name for b in Branch.query.all()}
 wb=Workbook(); ws=wb.active; ws.title=f'{month} 시재'; headers=['날짜','지점','구분','항목','결제수단','입금','출금','거래처/고객','메모','등록자']
 ws.append(headers)
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='14324A');cell.alignment=Alignment(horizontal='center')
 balance=0
 for x in items:
  signed=x.amount if x.direction=='입금' else -x.amount; balance+=signed if x.payment_method=='현금' else 0
  ws.append([x.ledger_date,branches.get(x.branch_id,'-'),x.direction,x.category,x.payment_method,x.amount if x.direction=='입금' else 0,x.amount if x.direction=='출금' else 0,x.counterparty or '',x.memo or '',x.created_by or ''])
 ws.append(['월 현금잔액','','','','',sum(x.amount for x in items if x.direction=='입금' and x.payment_method=='현금'),sum(x.amount for x in items if x.direction=='출금' and x.payment_method=='현금'),'','',''])
 for col,w in zip('ABCDEFGHIJ',[13,16,10,18,12,14,14,18,35,14]):ws.column_dimensions[col].width=w
 audit('시재 엑셀 다운로드','cash_ledger',month,f'지점 {branch_id or "전체"} · {len(items)}건');db.session.commit();out=io.BytesIO();wb.save(out);out.seek(0)
 return send_file(out,as_attachment=True,download_name=f'TrustFlow_{month}_시재관리.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.post('/cash-ledger/<int:ledger_id>/delete')
@login_required
def cash_ledger_delete(ledger_id):
 x=CashLedger.query.get_or_404(ledger_id); enforce_branch(x.branch_id)
 if x.reference_type=='card': flash('카드 자동매출은 카드매출 메뉴에서 취소 처리해주세요.','error')
 else: db.session.delete(x);db.session.commit();flash('시재 내역을 삭제했습니다.','success')
 return redirect(request.referrer or url_for('cash_ledger'))

@app.route('/card-sales',methods=['GET','POST'])
@login_required
def card_sales():
 prepare_database(); branch_id=scoped_branch_from_request()
 if request.method=='POST':
  if not is_admin():abort(403)
  bid=request.form.get('branch_id'); terminal=request.form.get('terminal_number','').strip()
  if not bid or not terminal:flash('지점과 카드단말기 등록번호를 입력해주세요.','error')
  elif CardTerminal.query.filter_by(terminal_number=terminal).first():flash('이미 등록된 단말기 번호입니다.','error')
  else:
   token=secrets.token_urlsafe(32);db.session.add(CardTerminal(branch_id=int(bid),provider=request.form.get('provider'),merchant_number=request.form.get('merchant_number'),terminal_number=terminal,api_token=token,memo=request.form.get('memo')));db.session.commit();flash(f'단말기 등록 완료 · 연동키: {token} (VAN사에 1회 전달)','success')
  return redirect(url_for('card_sales',branch_id=bid or ''))
 tq=CardTerminal.query
 txq=CardTransaction.query
 if branch_id:tq=tq.filter_by(branch_id=branch_id);txq=txq.filter_by(branch_id=branch_id)
 elif not is_admin():tq=tq.filter(CardTerminal.id==-1);txq=txq.filter(CardTransaction.id==-1)
 terminals=tq.order_by(CardTerminal.id.desc()).all(); transactions=txq.order_by(CardTransaction.paid_at.desc()).limit(300).all()
 return render_template('card_sales.html',terminals=terminals,transactions=transactions,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all(),branch_id=branch_id,branch_map={b.id:b for b in Branch.query.all()},terminal_map={t.id:t for t in terminals},today_total=sum(x.amount for x in transactions if x.status=='승인' and x.paid_at.date()==date.today()))

@app.post('/api/card-sales/<terminal_number>')
def card_sales_webhook(terminal_number):
 terminal=CardTerminal.query.filter_by(terminal_number=terminal_number,active=True).first_or_404(); token=request.headers.get('X-TrustFlow-Key') or request.args.get('token')
 if not secrets.compare_digest(token or '',terminal.api_token):abort(403)
 data=request.get_json(silent=True) or {}; approval=str(data.get('approval_number') or '').strip(); amount=abs(money(data.get('amount'))); status=str(data.get('status') or '승인')
 if not approval or amount<=0:return jsonify({'ok':False,'error':'approval_number and amount required'}),400
 existing=CardTransaction.query.filter_by(terminal_id=terminal.id,approval_number=approval).first()
 if existing:
  if existing.status==status:return jsonify({'ok':True,'duplicate':True})
  if status in ['취소','승인취소'] and existing.status=='승인':
   existing.status='취소';db.session.add(CashLedger(ledger_date=datetime.utcnow().date(),branch_id=terminal.branch_id,direction='출금',category='카드취소',amount=existing.amount,payment_method='카드',reference_type='card',reference_id=existing.id,counterparty=data.get('customer_name'),memo=f'승인취소 {approval}',created_by='카드단말기 자동수집'));db.session.commit();return jsonify({'ok':True,'cancelled':True,'transaction_id':existing.id})
  return jsonify({'ok':True,'duplicate':True})
 try:paid_at=datetime.fromisoformat(str(data.get('paid_at')).replace('Z','+00:00')).replace(tzinfo=None) if data.get('paid_at') else datetime.utcnow()
 except:paid_at=datetime.utcnow()
 tx=CardTransaction(terminal_id=terminal.id,branch_id=terminal.branch_id,approval_number=approval,paid_at=paid_at,amount=amount,card_company=data.get('card_company'),installment=str(data.get('installment') or ''),receipt_number=data.get('receipt_number'),status=status,raw_data=json.dumps(data,ensure_ascii=False))
 db.session.add(tx);db.session.flush();db.session.add(CashLedger(ledger_date=paid_at.date(),branch_id=terminal.branch_id,direction='입금' if status=='승인' else '출금',category='카드매출' if status=='승인' else '카드취소',amount=amount,payment_method='카드',reference_type='card',reference_id=tx.id,counterparty=data.get('customer_name'),memo=f'승인번호 {approval}',created_by='카드단말기 자동수집'));db.session.commit()
 return jsonify({'ok':True,'transaction_id':tx.id})

@app.route('/ob-management')
@login_required
def ob_management():
 cutoff=date.today()-timedelta(days=548); sq=apply_branch_scope(Sale.query,Sale).filter(Sale.opening_date<=cutoff).order_by(Sale.opening_date.asc()).all(); latest={}
 for s in sq:
  key=normalize_phone(s.customer_phone)
  if key and (key not in latest or s.opening_date>latest[key].opening_date):latest[key]=s
 customers={normalize_phone(c.phone):c for c in Customer.query.filter(Customer.phone.in_(list(latest.keys()) or ['__none__'])).all()}; logs=ContactLog.query
 if not is_admin():logs=logs.filter_by(branch_id=current_branch_id())
 last_logs={}
 for x in logs.order_by(ContactLog.contacted_at.desc()).all():last_logs.setdefault(x.customer_id,x)
 return render_template('ob_management.html',rows=[(customers.get(p),s) for p,s in latest.items() if customers.get(p)],last_logs=last_logs,cutoff=cutoff,branches={b.id:b for b in Branch.query.all()})

@app.post('/customers/<int:cid>/contact-log')
@login_required
def contact_log_add(cid):
 c=Customer.query.get_or_404(cid)
 if not customer_allowed(c):abort(403)
 sale=Sale.query.filter_by(customer_phone=c.phone).order_by(Sale.opening_date.desc()).first(); bid=current_branch_id() if not is_admin() else (request.form.get('branch_id') or c.branch_id or (sale.branch_id if sale else None))
 if not bid:flash('담당 지점을 확인할 수 없습니다.','error');return redirect(url_for('customer_detail',cid=cid))
 enforce_branch(bid); note=request.form.get('note','').strip()
 if not note:flash('통화내용을 입력해주세요.','error');return redirect(url_for('customer_detail',cid=cid))
 db.session.add(ContactLog(customer_id=cid,branch_id=int(bid),staff_name=session.get('display_name') or session.get('username'),channel=request.form.get('channel','전화'),outcome=request.form.get('outcome','상담완료'),note=note,next_contact_date=parse_date(request.form.get('next_contact_date'))));db.session.commit();flash('상담 기록을 저장했습니다.','success');return redirect(url_for('customer_detail',cid=cid))

@app.route('/legal-cases',methods=['GET','POST'])
@login_required
def legal_cases():
 prepare_database(); branch_id=scoped_branch_from_request()
 if request.method=='POST':
  bid=current_branch_id() if not is_admin() else request.form.get('branch_id'); cid=request.form.get('customer_id')
  if not bid or not cid:flash('고객과 담당지점을 선택해주세요.','error')
  else:
   db.session.add(LegalCase(customer_id=int(cid),branch_id=int(bid),case_type=request.form.get('case_type','환수'),claim_amount=money(request.form.get('claim_amount')),incident_date=parse_date(request.form.get('incident_date')),reason=request.form.get('reason'),evidence=request.form.get('evidence'),debtor_address=request.form.get('debtor_address'),demand_due_date=parse_date(request.form.get('demand_due_date')),status='자료수집',assigned_staff=request.form.get('assigned_staff') or session.get('display_name'),created_by=session.get('display_name') or session.get('username')));db.session.commit();flash('환수 법률업무가 등록되었습니다.','success')
  return redirect(url_for('legal_cases',branch_id=bid or ''))
 q=LegalCase.query
 if branch_id:q=q.filter_by(branch_id=branch_id)
 elif not is_admin():q=q.filter(LegalCase.id==-1)
 items=q.order_by(LegalCase.created_at.desc()).all(); customer_map={c.id:c for c in Customer.query.filter(Customer.id.in_([x.customer_id for x in items] or [0])).all()}
 allowed_sales=apply_branch_scope(Sale.query,Sale).order_by(Sale.customer_name).all(); phones=list(dict.fromkeys([s.customer_phone for s in allowed_sales if s.customer_phone])); customers=Customer.query.filter(Customer.phone.in_(phones or ['__none__'])).order_by(Customer.name).all()
 return render_template('legal_cases.html',items=items,customer_map=customer_map,customers=customers,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all(),branch_map={b.id:b for b in Branch.query.all()},branch_id=branch_id,today=date.today().isoformat())

@app.post('/legal-cases/<int:case_id>/status')
@login_required
def legal_case_status(case_id):
 x=LegalCase.query.get_or_404(case_id);enforce_branch(x.branch_id);x.status=request.form.get('status',x.status);x.memo=request.form.get('memo',x.memo);db.session.commit();flash('법률업무 상태를 변경했습니다.','success');return redirect(url_for('legal_cases'))

@app.get('/legal-cases/<int:case_id>/notice')
@login_required
def legal_case_notice(case_id):
 x=LegalCase.query.get_or_404(case_id);enforce_branch(x.branch_id);c=Customer.query.get_or_404(x.customer_id);b=Branch.query.get(x.branch_id)
 text_body=f'''내용증명\n\n수신인: {c.name}\n주소: {x.debtor_address or c.address_road or c.address_jibun or '[주소 확인 필요]'}\n발신인: {b.name if b else 'TrustFlow 등록 사업자'}\n\n제목: {x.case_type} 관련 금원 지급 요청\n\n1. 발생일: {x.incident_date or '[확인 필요]'}\n2. 청구금액: {x.claim_amount:,}원\n3. 청구사유: {x.reason or '[구체적 사실관계 입력 필요]'}\n4. 보유 증빙: {x.evidence or '[계약서·입금내역·대화내역 등 확인 필요]'}\n5. 지급기한: {x.demand_due_date or '[기한 입력 필요]'}\n\n위 기한까지 지급 또는 협의가 없을 경우 지급명령·소액사건심판 등 적법한 절차를 검토할 수 있음을 알려드립니다.\n\n작성일: {date.today()}\n발신인: ____________________\n\n※ 본 문서는 내부 업무용 초안입니다. 발송 전 사실관계·계약·개인정보·관할법원을 확인하고 필요한 경우 변호사 또는 법률구조기관의 검토를 받으세요.'''
 audit('법률서식 다운로드','legal_case',x.id,f'{c.name} · {x.case_type}',x.branch_id);db.session.commit();out=io.BytesIO(text_body.encode('utf-8-sig'));return send_file(out,as_attachment=True,download_name=f'{c.name}_내용증명_초안.txt',mimetype='text/plain; charset=utf-8')

@app.route('/')
@login_required
def dashboard():
 prepare_database(); today=date.today(); selected=parse_date(request.args.get('date')) or today
 start=date(today.year,today.month,1); end=add_months(start,1)
 tq=task_query_scoped()
 tasks=tq.filter(CustomerTask.due_date==selected).order_by(CustomerTask.status.asc(),CustomerTask.id.desc()).all()
 overdue=task_query_scoped().filter(CustomerTask.due_date<today,CustomerTask.status.in_(['처리예정','연락안됨','연기'])).order_by(CustomerTask.due_date.asc()).limit(50).all()
 month_tasks=task_query_scoped().filter(CustomerTask.due_date>=start,CustomerTask.due_date<end).all()
 counts={}
 for t in month_tasks: counts[t.due_date.isoformat()]=counts.get(t.due_date.isoformat(),0)+1
 sq=apply_branch_scope(Sale.query,Sale)
 today_sale_items=sq.filter(Sale.opening_date==today).order_by(Sale.id.desc()).all()
 pq=payback_query_scoped()
 pending_paybacks=pq.filter(Payback.status!='완료').count()
 today_paybacks=payback_query_scoped().filter(Payback.due_date==today,Payback.status!='완료').all()
 branches={b.id:b for b in Branch.query.all()}
 sales_map={s.id:s for s in Sale.query.filter(Sale.id.in_([p.sale_id for p in today_paybacks] or [0])).all()}
 cal=calendar.Calendar(firstweekday=6); weeks=cal.monthdayscalendar(today.year,today.month)
 return render_template('dashboard.html',today=today,selected=selected,tasks=tasks,overdue=overdue,counts=counts,weeks=weeks,year=today.year,month=today.month,today_sales=len(today_sale_items),today_sale_items=today_sale_items,pending_paybacks=pending_paybacks,today_paybacks=today_paybacks,sales_map=sales_map,branches=branches)

@app.get('/notifications')
@login_required
def notifications():
 prepare_database(); today=date.today(); open_states=['처리예정','연락안됨','연기']
 overdue_tasks=task_query_scoped().filter(CustomerTask.due_date<today,CustomerTask.status.in_(open_states)).order_by(CustomerTask.due_date.asc()).limit(100).all()
 today_tasks=task_query_scoped().filter(CustomerTask.due_date==today,CustomerTask.status.in_(open_states)).order_by(CustomerTask.id.desc()).limit(100).all()
 pq=payback_query_scoped(); due_paybacks=pq.filter(Payback.status!='완료',Payback.due_date<=today).order_by(Payback.due_date.asc()).limit(100).all()
 pending_approvals=pq.filter(Payback.status!='완료',Payback.approval_status=='승인대기').order_by(Payback.due_date.asc()).limit(100).all() if can_approve_payback() else []
 lq=apply_branch_scope(LegalCase.query,LegalCase)
 legal_deadlines=lq.filter(LegalCase.status.notin_(['완료','종결']),LegalCase.demand_due_date.isnot(None),LegalCase.demand_due_date<=today+timedelta(days=3)).order_by(LegalCase.demand_due_date.asc()).limit(100).all()
 sale_ids=list({p.sale_id for p in due_paybacks+pending_approvals}); sales_map={s.id:s for s in Sale.query.filter(Sale.id.in_(sale_ids or [0])).all()}
 customer_ids=list({x.customer_id for x in legal_deadlines if x.customer_id}); customer_map={c.id:c for c in Customer.query.filter(Customer.id.in_(customer_ids or [0])).all()}
 return render_template('notifications.html',today=today,overdue_tasks=overdue_tasks,today_tasks=today_tasks,due_paybacks=due_paybacks,pending_approvals=pending_approvals,legal_deadlines=legal_deadlines,sales_map=sales_map,customer_map=customer_map)

@app.post('/tasks/<int:task_id>/status')
@login_required
def task_status(task_id):
 t=CustomerTask.query.get_or_404(task_id);
 if t.sale_id:
  enforce_branch(Sale.query.get_or_404(t.sale_id).branch_id)
 status=request.form.get('status','처리예정'); old_due=t.due_date; t.status=status; t.result_memo=request.form.get('memo','').strip() or t.result_memo
 if status=='완료':
  t.completed_at=datetime.utcnow(); t.completed_by=session.get('display_name') or session.get('username')
 else:
  t.completed_at=None; t.completed_by=None
 if status=='연기' and request.form.get('due_date'):
  new_due=parse_date(request.form.get('due_date')) or t.due_date
  if new_due!=old_due:
   history=f'[연기 {datetime.now().strftime("%Y-%m-%d %H:%M")}] {old_due} → {new_due} / {session.get("display_name") or session.get("username")}'
   t.result_memo=(t.result_memo+'\\n' if t.result_memo else '')+history
  t.due_date=new_due
 db.session.commit();flash('고객약속 상태가 변경되었습니다.','success');return redirect(request.referrer or url_for('dashboard'))

@app.route('/tasks/<int:task_id>/edit',methods=['GET','POST'])
@login_required
def task_edit(task_id):
 t=CustomerTask.query.get_or_404(task_id)
 if t.sale_id: enforce_branch(Sale.query.get_or_404(t.sale_id).branch_id)
 if request.method=='POST':
  old_due=t.due_date; t.task_type=request.form.get('task_type','기타'); t.title=request.form.get('title','').strip() or t.title
  t.description=request.form.get('description'); t.due_date=parse_date(request.form.get('due_date')) or t.due_date; t.assigned_staff=request.form.get('assigned_staff'); t.status=request.form.get('status','처리예정'); t.result_memo=request.form.get('result_memo')
  if t.due_date!=old_due:
   log=f'[날짜수정 {datetime.now().strftime("%Y-%m-%d %H:%M")}] {old_due} → {t.due_date} / {session.get("display_name") or session.get("username")}'
   t.result_memo=(t.result_memo+'\\n' if t.result_memo else '')+log
  if t.status=='완료':
   if not t.completed_at:t.completed_at=datetime.utcnow()
   t.completed_by=session.get('display_name') or session.get('username')
  else:
   t.completed_at=None; t.completed_by=None
  if t.task_type=='페이백 지급' and t.sale_id:
   p=Payback.query.filter_by(sale_id=t.sale_id).first()
   if p: p.due_date=t.due_date; p.status='완료' if t.status=='완료' else p.status
  db.session.commit(); flash('고객약속이 수정되었습니다.','success'); return redirect(url_for('dashboard',date=t.due_date.isoformat()))
 return render_template('task_edit.html',t=t,staff=(User.query.filter_by(active=True,branch_id=current_branch_id()).order_by(User.display_name).all() if not is_admin() else User.query.filter_by(active=True).order_by(User.display_name,User.username).all()))

@app.route('/customers')
@login_required
def customers():
 prepare_database(); q=request.args.get('q','').strip(); month=request.args.get('month','').strip(); branch_id=request.args.get('branch_id','').strip()
 if not is_admin(): branch_id=str(current_branch_id() or '')
 query=customer_query_scoped()
 if q:
  phone_q=normalize_phone(q); query=query.filter(or_(Customer.name.ilike(f'%{q}%'),Customer.phone.ilike(f'%{phone_q or q}%'),Customer.address_road.ilike(f'%{q}%'),Customer.address_jibun.ilike(f'%{q}%'),Customer.address_detail.ilike(f'%{q}%')))
 sale_scope=Sale.query
 if branch_id:
  try:sale_scope=sale_scope.filter(Sale.branch_id==int(branch_id))
  except:pass
 if month:
  try:
   y,m=map(int,month.split('-')); mstart=date(y,m,1); mend=add_months(mstart,1)
   sale_scope=sale_scope.filter(Sale.opening_date>=mstart,Sale.opening_date<mend)
  except: pass
 phones=[r[0] for r in sale_scope.with_entities(Sale.customer_phone).distinct().all() if r[0]]
 if branch_id or month: query=query.filter(Customer.phone.in_(phones or ['__none__']))
 customers_list=query.order_by(Customer.created_at.desc()).all()
 sale_map={}
 for c in customers_list:
  sq=Sale.query.filter_by(customer_phone=c.phone) if c.phone else Sale.query.filter(Sale.id==-1)
  if branch_id:
   try:sq=sq.filter(Sale.branch_id==int(branch_id))
   except:pass
  if month:
   try:sq=sq.filter(Sale.opening_date>=mstart,Sale.opening_date<mend)
   except:pass
  sale_map[c.id]=sq.order_by(Sale.opening_date.desc(),Sale.id.desc()).first()
 months=[r[0].strftime('%Y-%m') for r in db.session.query(Sale.opening_date).filter(Sale.opening_date.isnot(None)).order_by(Sale.opening_date.desc()).all()]
 months=list(dict.fromkeys(months))
 return render_template('customers.html',customers=customers_list,q=q,month=month,months=months,branches=Branch.query.filter_by(active=True).all(),branch_id=branch_id,sale_map=sale_map)

@app.route('/customers/<int:cid>')
@login_required
def customer_detail(cid):
 c=Customer.query.get_or_404(cid)
 if not customer_allowed(c):abort(403)
 sq=Sale.query.filter_by(customer_phone=c.phone) if c.phone else Sale.query.filter(Sale.id==-1)
 if not is_admin(): sq=sq.filter(Sale.branch_id==current_branch_id())
 sale_history=sq.order_by(Sale.opening_date.desc(),Sale.id.desc()).all()
 if not sale_history and not is_admin(): abort(403)
 sale_ids=[s.id for s in sale_history]
 tasks=CustomerTask.query.filter(CustomerTask.sale_id.in_(sale_ids or [0])).order_by(CustomerTask.due_date.desc()).all()
 open_tasks=[t for t in tasks if t.status not in ['완료','취소']]
 paybacks=Payback.query.filter(Payback.sale_id.in_(sale_ids or [0])).order_by(Payback.due_date.desc()).all()
 doc_counts=dict(db.session.query(SaleDocument.sale_id,db.func.count(SaleDocument.id)).filter(SaleDocument.sale_id.in_(sale_ids or [0])).group_by(SaleDocument.sale_id).all())
 branches={b.id:b for b in Branch.query.all()}
 household=[]
 if c.address_key:
  hq=customer_query_scoped().filter(Customer.address_key==c.address_key,Customer.id!=c.id)
  if not is_admin():
   allowed_phones=[r[0] for r in Sale.query.filter_by(branch_id=current_branch_id()).with_entities(Sale.customer_phone).distinct().all() if r[0]]; hq=hq.filter(Customer.phone.in_(allowed_phones or ['__none__']))
  household=hq.order_by(Customer.name).all()
 contact_q=ContactLog.query.filter_by(customer_id=c.id)
 if not is_admin():contact_q=contact_q.filter_by(branch_id=current_branch_id())
 contact_logs=contact_q.order_by(ContactLog.contacted_at.desc()).limit(100).all()
 return render_template('customer_detail.html',customer=c,sales=sale_history,tasks=tasks,open_tasks=open_tasks,paybacks=paybacks,doc_counts=doc_counts,branches=branches,household=household,contact_logs=contact_logs)

@app.route('/customers/new',methods=['GET','POST'])
@login_required
def customer_new():
 if request.method=='POST':
  name=request.form.get('name','').strip()
  if not name:flash('고객명을 입력해주세요.','error');return redirect(url_for('customer_new'))
  road=request.form.get('address_road','').strip(); jibun=request.form.get('address_jibun','').strip(); detail=request.form.get('address_detail','').strip(); key=request.form.get('address_key','').strip() or ('|'.join([road,jibun,detail]).lower().replace(' ',''))
  bid=current_branch_id() if not is_admin() else request.form.get('branch_id')
  if not bid:flash('고객을 등록할 지점을 선택해주세요.','error');return redirect(url_for('customer_new'))
  enforce_branch(bid)
  db.session.add(Customer(name=name,phone=normalize_phone(request.form.get('phone','')),carrier=request.form.get('carrier',''),status=request.form.get('status','상담중'),memo=request.form.get('memo',''),address_road=road,address_jibun=jibun,address_detail=detail,address_key=key,branch_id=int(bid),company_code=current_company()));db.session.commit();flash('고객이 등록되었습니다.','success');return redirect(url_for('customers'))
 return render_template('customer_form.html',customer=None,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all())
@app.route('/customers/<int:cid>/edit',methods=['GET','POST'])
@login_required
def customer_edit(cid):
 c=Customer.query.get_or_404(cid)
 if not customer_allowed(c):abort(403)
 if request.method=='POST':
  road=request.form.get('address_road','').strip(); jibun=request.form.get('address_jibun','').strip(); detail=request.form.get('address_detail','').strip(); key=request.form.get('address_key','').strip() or ('|'.join([road,jibun,detail]).lower().replace(' ',''))
  bid=current_branch_id() if not is_admin() else (request.form.get('branch_id') or c.branch_id)
  enforce_branch(bid)
  c.name=request.form.get('name','').strip();c.phone=normalize_phone(request.form.get('phone',''));c.carrier=request.form.get('carrier','');c.status=request.form.get('status','상담중');c.memo=request.form.get('memo','');c.address_road=road;c.address_jibun=jibun;c.address_detail=detail;c.address_key=key;c.branch_id=int(bid);db.session.commit();flash('고객정보가 수정되었습니다.','success');return redirect(url_for('customers'))
 return render_template('customer_form.html',customer=c,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all())

@app.post('/customers/<int:cid>/delete')
@login_required
@admin_required
def customer_delete(cid):
 c=Customer.query.get_or_404(cid)
 if not customer_allowed(c):abort(403)
 if c.phone and Sale.query.filter_by(customer_phone=c.phone).count():
  flash('개통이력이 있는 고객은 삭제할 수 없습니다. 고객정보 수정으로 관리해주세요.','error'); return redirect(url_for('customers'))
 CustomerTask.query.filter_by(customer_id=c.id).delete(synchronize_session=False); db.session.delete(c); db.session.commit()
 flash('고객이 삭제되었습니다.','success'); return redirect(url_for('customers'))

@app.route('/inventory')
@login_required
def inventory():
 q=request.args.get('q','').strip(); status=request.args.get('status','').strip(); branch_id=request.args.get('branch_id','').strip()
 query=Inventory.query
 if not is_admin():
  branch_id=str(current_branch_id() or ''); query=query.filter(Inventory.branch_id==current_branch_id()) if current_branch_id() else query.filter(Inventory.id==-1)
 elif branch_id:
  try:query=query.filter(Inventory.branch_id==int(branch_id))
  except:pass
 if q: query=query.filter(or_(Inventory.serial_number.ilike(f'%{q}%'),Inventory.model.ilike(f'%{q}%'),Inventory.color.ilike(f'%{q}%')))
 if status: query=query.filter_by(status=status)
 items=query.order_by(Inventory.received_date.desc(),Inventory.id.desc()).all()
 partners={p.id:p for p in Partner.query.all()}; branches={b.id:b for b in Branch.query.all()}; today=date.today()
 return render_template('inventory.html',items=items,partners=partners,branches=branches,branch_list=Branch.query.filter_by(active=True).order_by(Branch.id).all(),q=q,status=status,branch_id=branch_id,today=today)

@app.route('/inventory/new',methods=['GET','POST'])
@login_required
def inventory_new():
 partners=Partner.query.filter_by(active=True).order_by(Partner.name).all()
 branches=Branch.query.filter_by(active=True).order_by(Branch.name).all()
 device_models=DeviceMaster.query.filter_by(active=True).order_by(DeviceMaster.manufacturer,DeviceMaster.sort_order,DeviceMaster.model).all(); device_master_data=[{'manufacturer':d.manufacturer,'model':d.model,'capacities':d.capacities or '','colors':d.colors or ''} for d in device_models]
 if request.method=='POST':
  serial=request.form.get('serial_number','').strip()
  if not serial or not request.form.get('model','').strip():flash('일련번호와 모델명은 필수입니다.','error');return redirect(url_for('inventory_new'))
  if Inventory.query.filter_by(serial_number=serial).first():flash('이미 등록된 일련번호입니다.','error');return redirect(url_for('inventory_new'))
  branch_id=request.form.get('branch_id') or None
  if not is_admin(): branch_id=current_branch_id()
  if not branch_id: flash('보유지점을 지정해주세요.','error'); return redirect(url_for('inventory_new'))
  item=Inventory(serial_number=serial,partner_id=request.form.get('partner_id') or None,carrier=request.form.get('carrier'),manufacturer=request.form.get('manufacturer'),model=request.form.get('model').strip(),capacity=request.form.get('capacity'),color=request.form.get('color'),received_date=parse_date(request.form.get('received_date')) or date.today(),purchase_price=money(request.form.get('purchase_price')),storage_location=request.form.get('storage_location'),branch_id=branch_id,status='보유중',memo=request.form.get('memo'))
  db.session.add(item);db.session.flush()
  db.session.add(InventoryMovement(inventory_id=item.id,action='입고등록',to_branch_id=item.branch_id,to_status='보유중',processed_by=session.get('display_name') or session.get('username'),memo=item.memo))
  db.session.commit();flash('단말기가 입고 등록되었습니다.','success');return redirect(url_for('inventory'))
 return render_template('inventory_form.html',partners=partners,branches=branches,today=date.today().isoformat(),item=None,device_models=device_models,device_master_data=device_master_data)

@app.get('/api/inventory/<path:serial>')
@login_required
def inventory_lookup(serial):
 i=Inventory.query.filter_by(serial_number=serial).first()
 if not i:return jsonify({'ok':False}),404
 enforce_branch(i.branch_id)
 p=Partner.query.get(i.partner_id) if i.partner_id else None
 b=Branch.query.get(i.branch_id) if i.branch_id else None
 return jsonify({'ok':True,'id':i.id,'serial_number':i.serial_number,'model':i.model,'capacity':i.capacity or '','color':i.color or '','carrier':i.carrier or '','manufacturer':i.manufacturer or '','partner_id':i.partner_id,'partner':p.name if p else '','branch_id':i.branch_id,'branch':b.name if b else '','status':i.status})


@app.route('/inventory/<int:iid>')
@login_required
def inventory_detail(iid):
 item=Inventory.query.get_or_404(iid)
 enforce_branch(item.branch_id)
 partner=Partner.query.get(item.partner_id) if item.partner_id else None
 branch=Branch.query.get(item.branch_id) if item.branch_id else None
 moves=InventoryMovement.query.filter_by(inventory_id=item.id).order_by(InventoryMovement.created_at.desc(),InventoryMovement.id.desc()).all()
 branch_map={b.id:b for b in Branch.query.all()}
 sale=Sale.query.get(item.sale_id) if item.sale_id else None
 return render_template('inventory_detail.html',item=item,partner=partner,branch=branch,moves=moves,branch_map=branch_map,sale=sale)

@app.route('/inventory/<int:iid>/edit',methods=['GET','POST'])
@login_required
def inventory_edit(iid):
 item=Inventory.query.get_or_404(iid)
 enforce_branch(item.branch_id)
 if request.method=='POST':
  old_branch=item.branch_id; old_status=item.status
  item.partner_id=request.form.get('partner_id') or None; item.carrier=request.form.get('carrier'); item.manufacturer=request.form.get('manufacturer')
  item.model=request.form.get('model','').strip() or item.model; item.capacity=request.form.get('capacity'); item.color=request.form.get('color')
  item.received_date=parse_date(request.form.get('received_date')) or item.received_date; item.purchase_price=money(request.form.get('purchase_price'))
  item.storage_location=request.form.get('storage_location'); item.memo=request.form.get('memo')
  new_branch=request.form.get('branch_id')
  if item.status!='판매완료': item.branch_id=int(new_branch) if new_branch else None
  db.session.add(InventoryMovement(inventory_id=item.id,action='정보수정',from_branch_id=old_branch,to_branch_id=item.branch_id,from_status=old_status,to_status=item.status,processed_by=session.get('display_name') or session.get('username'),memo='재고정보 수정'))
  db.session.commit(); flash('재고정보가 수정되었습니다.','success'); return redirect(url_for('inventory_detail',iid=item.id))
 device_models=DeviceMaster.query.filter_by(active=True).order_by(DeviceMaster.manufacturer,DeviceMaster.sort_order,DeviceMaster.model).all()
 device_master_data=[{'manufacturer':d.manufacturer,'model':d.model,'capacities':d.capacities or '','colors':d.colors or ''} for d in device_models]
 return render_template('inventory_form.html',partners=Partner.query.filter_by(active=True).order_by(Partner.name).all(),branches=Branch.query.filter_by(active=True).order_by(Branch.id).all(),today=item.received_date.isoformat(),item=item,device_models=device_models,device_master_data=device_master_data)

@app.route('/inventory/<int:iid>/move',methods=['GET','POST'])
@login_required
def inventory_move(iid):
 item=Inventory.query.get_or_404(iid)
 enforce_branch(item.branch_id)
 if not is_admin(): abort(403)
 if item.status in ['판매완료','거래처반납','반품']:
  flash('판매완료/반품 처리된 단말기는 지점이동할 수 없습니다.','error'); return redirect(url_for('inventory_detail',iid=iid))
 if request.method=='POST':
  to_branch=request.form.get('to_branch_id')
  if not to_branch: flash('이동할 지점을 선택해주세요.','error'); return redirect(url_for('inventory_move',iid=iid))
  to_branch=int(to_branch)
  if item.branch_id==to_branch: flash('현재 보유지점과 동일합니다.','error'); return redirect(url_for('inventory_move',iid=iid))
  old=item.branch_id; old_status=item.status
  item.branch_id=to_branch; item.status='보유중'
  db.session.add(InventoryMovement(inventory_id=item.id,action='지점이동',from_branch_id=old,to_branch_id=to_branch,from_status=old_status,to_status='보유중',processed_by=session.get('display_name') or session.get('username'),memo=request.form.get('memo')))
  db.session.commit(); flash('지점이동이 완료되었습니다.','success'); return redirect(url_for('inventory_detail',iid=iid))
 return render_template('inventory_move.html',item=item,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all(),current_branch=Branch.query.get(item.branch_id) if item.branch_id else None,mode='move')

@app.post('/inventory/<int:iid>/action')
@login_required
def inventory_action(iid):
 item=Inventory.query.get_or_404(iid); enforce_branch(item.branch_id); action=request.form.get('action',''); old_status=item.status; old_branch=item.branch_id
 if item.status=='판매완료':
  flash('판매완료 단말기는 판매일보에서 개통건을 수정/삭제한 뒤 처리해주세요.','error'); return redirect(url_for('inventory_detail',iid=iid))
 if action=='예약':
  item.status='예약'
 elif action=='예약취소':
  item.status='보유중'
 elif action=='회수':
  to_branch=request.form.get('to_branch_id')
  if not to_branch: flash('회수할 지점을 선택해주세요.','error'); return redirect(url_for('inventory_detail',iid=iid))
  item.branch_id=int(to_branch); item.status='보유중'
 elif action=='거래처반납':
  item.status='거래처반납'
 elif action=='반품':
  item.status='반품'
 else:
  abort(400)
 db.session.add(InventoryMovement(inventory_id=item.id,action=action,from_branch_id=old_branch,to_branch_id=item.branch_id,from_status=old_status,to_status=item.status,processed_by=session.get('display_name') or session.get('username'),memo=request.form.get('memo')))
 db.session.commit(); flash(f'{action} 처리가 완료되었습니다.','success'); return redirect(url_for('inventory_detail',iid=iid))

@app.post('/inventory/<int:iid>/delete')
@login_required
@admin_required
def inventory_delete(iid):
 item=Inventory.query.get_or_404(iid)
 if item.sale_id or item.status=='판매완료':
  flash('판매 이력이 있는 재고는 삭제할 수 없습니다.','error'); return redirect(url_for('inventory_detail',iid=iid))
 movement_count=InventoryMovement.query.filter_by(inventory_id=item.id).count()
 if movement_count>1:
  flash('이동/상태변경 이력이 있는 재고는 삭제할 수 없습니다. 반품 또는 거래처반납으로 처리해주세요.','error'); return redirect(url_for('inventory_detail',iid=iid))
 InventoryMovement.query.filter_by(inventory_id=item.id).delete(synchronize_session=False); db.session.delete(item); db.session.commit()
 flash('잘못 등록된 재고가 삭제되었습니다.','success'); return redirect(url_for('inventory'))

@app.route('/partners')
@login_required
def partners():
 plist=Partner.query.order_by(Partner.name).all();rows=[]
 for p in plist:
  stock=Inventory.query.filter_by(partner_id=p.id,status='보유중').count();monthly=Sale.query.filter(Sale.partner_id==p.id,Sale.opening_date>=date.today().replace(day=1)).all();unsettled=sum(1 for s in monthly if s.status!='정산완료');rows.append((p,stock,len(monthly),sum(s.settlement_amount_v2 or money(s.settlement) for s in monthly),unsettled))
 return render_template('partners.html',rows=rows)
@app.route('/partners/new',methods=['GET','POST'])
@login_required
def partner_new():
 if request.method=='POST':
  name=request.form.get('name','').strip()
  if not name:flash('거래처명을 입력해주세요.','error');return redirect(url_for('partner_new'))
  if Partner.query.filter_by(name=name).first():flash('이미 등록된 거래처입니다.','error');return redirect(url_for('partner_new'))
  db.session.add(Partner(name=name,category=request.form.get('category'),contact_name=request.form.get('contact_name'),phone=request.form.get('phone'),settlement_cycle=request.form.get('settlement_cycle'),memo=request.form.get('memo')));db.session.commit();flash('거래처가 등록되었습니다.','success');return redirect(url_for('partners'))
 return render_template('partner_form.html')


@app.route('/branches')
@login_required
def branches():
 rows=[]
 for b in Branch.query.order_by(Branch.id).all():
  rows.append((b,Inventory.query.filter_by(branch_id=b.id,status='보유중').count(),Sale.query.filter_by(branch_id=b.id).count(),User.query.filter_by(branch_id=b.id).count()))
 return render_template('branches.html',rows=rows)

@app.route('/branches/new',methods=['GET','POST'])
@login_required
@admin_required
def branch_new():
 if request.method=='POST':
  name=request.form.get('name','').strip()
  if not name: flash('지점명을 입력해주세요.','error'); return redirect(url_for('branch_new'))
  db.session.add(Branch(name=name,code=request.form.get('code'),address=request.form.get('address'),phone=request.form.get('phone'),manager_name=request.form.get('manager_name'),memo=request.form.get('memo'),active=True,company_code=current_company())); db.session.commit()
  flash('지점이 등록되었습니다.','success'); return redirect(url_for('branches'))
 return render_template('branch_form.html',branch=None)

@app.route('/branches/<int:bid>/edit',methods=['GET','POST'])
@login_required
@admin_required
def branch_edit(bid):
 b=Branch.query.get_or_404(bid)
 if request.method=='POST':
  b.name=request.form.get('name','').strip(); b.code=request.form.get('code'); b.address=request.form.get('address'); b.phone=request.form.get('phone'); b.manager_name=request.form.get('manager_name'); b.memo=request.form.get('memo'); b.active=request.form.get('active')=='1'
  db.session.commit(); flash('지점정보가 수정되었습니다.','success'); return redirect(url_for('branches'))
 return render_template('branch_form.html',branch=b)

@app.route('/partners/<int:pid>/edit',methods=['GET','POST'])
@login_required
def partner_edit(pid):
 p=Partner.query.get_or_404(pid)
 if request.method=='POST':
  p.name=request.form.get('name','').strip(); p.category=request.form.get('category'); p.contact_name=request.form.get('contact_name'); p.phone=request.form.get('phone'); p.settlement_cycle=request.form.get('settlement_cycle'); p.memo=request.form.get('memo'); p.active=request.form.get('active')=='1'
  db.session.commit(); flash('거래처가 수정되었습니다.','success'); return redirect(url_for('partners'))
 return render_template('partner_form.html',partner=p)

@app.post('/partners/<int:pid>/delete')
@login_required
@admin_required
def partner_delete(pid):
 p=Partner.query.get_or_404(pid)
 if Inventory.query.filter_by(partner_id=p.id).count() or Sale.query.filter_by(partner_id=p.id).count():
  flash('재고/판매 이력이 있는 거래처는 삭제할 수 없습니다. 비활성으로 변경해주세요.','error'); return redirect(url_for('partners'))
 db.session.delete(p); db.session.commit(); flash('거래처가 삭제되었습니다.','success'); return redirect(url_for('partners'))

@app.route('/sales/new',methods=['GET','POST'])
@login_required
def sale_new():
 prepare_database(); staff=(User.query.filter_by(active=True,branch_id=current_branch_id()).order_by(User.display_name).all() if not is_admin() else User.query.filter_by(active=True).order_by(User.display_name,User.username).all());partners=Partner.query.filter_by(active=True).order_by(Partner.name).all();branches=(Branch.query.filter_by(id=current_branch_id()).all() if not is_admin() else Branch.query.filter_by(active=True).order_by(Branch.id).all());plans=PlanMaster.query.filter_by(active=True).order_by(PlanMaster.carrier,PlanMaster.sort_order,PlanMaster.name).all(); plan_data=[{'carrier':p.carrier,'name':p.name} for p in plans]
 if request.method=='POST':
  name=request.form.get('customer_name','').strip(); opening=parse_date(request.form.get('opening_date')) or date.today()
  if not name:flash('고객명을 입력해주세요.','error');return redirect(url_for('sale_new'))
  phone=request.form.get('customer_phone','').strip(); customer=Customer.query.filter_by(phone=phone).first() if phone else None
  if not customer:customer=Customer(name=name,phone=phone,carrier=request.form.get('carrier'),status='개통고객');db.session.add(customer);db.session.flush()
  serial=request.form.get('serial_number','').strip(); inv=Inventory.query.filter_by(serial_number=serial).first() if serial else None
  rebate=money(request.form.get('rebate'));verbal=money(request.form.get('verbal_extra'));deduct=money(request.form.get('deduction'));support=money(request.form.get('extra_support'));payback=money(request.form.get('customer_payback'));opening_type=request.form.get('opening_type');sim_type=request.form.get('sim_payment_type','없음');settlement,tax,margin,transfer_fee=calc_settlement(rebate,verbal,deduct,support,payback,opening_type,sim_type)
  plan_due=opening+timedelta(days=183) if request.form.get('next_plan','').strip() else None
  internet_due=parse_date(request.form.get('internet_cancel_due_date'));payback_due=parse_date(request.form.get('payback_due_date'))
  sale=Sale(customer_name=name,customer_phone=phone,customer_birth=request.form.get('customer_birth'),opening_date=opening,carrier=request.form.get('carrier'),opening_type=opening_type,status='개통완료',manufacturer=(inv.manufacturer if inv else request.form.get('manufacturer')),device=(inv.model if inv else request.form.get('device')),color=(inv.color if inv else request.form.get('color')),storage=(inv.capacity if inv else request.form.get('storage')),serial_number=serial,plan=request.form.get('current_plan'),current_plan=request.form.get('current_plan'),next_plan=request.form.get('next_plan'),plan_change_due_date=plan_due,partner_id=(inv.partner_id if inv else (request.form.get('partner_id') or None)),inventory_id=(inv.id if inv else None),visit_source=request.form.get('visit_source'),branch_id=(current_branch_id() if not is_admin() else (request.form.get('branch_id') or None)),assigned_staff=request.form.get('assigned_staff') or session.get('display_name') or session.get('username'),created_by=session.get('display_name') or session.get('username'),rebate=rebate,verbal_extra=verbal,deduction=deduct,extra_support=support,settlement_amount_v2=settlement,tax_rate=.133,tax_amount=tax,customer_payback=payback,transfer_fee=transfer_fee,sim_payment_type=sim_type,sim_fee=7700,final_margin=margin,settlement=str(settlement),margin=str(margin),internet_carrier=request.form.get('internet_carrier'),internet_subscriber=request.form.get('internet_subscriber'),internet_install_date=parse_date(request.form.get('internet_install_date')),internet_cancel_due_date=internet_due,payback_due_date=payback_due,memo=request.form.get('memo'))
  db.session.add(sale);db.session.flush()
  if inv:inv.status='판매완료';inv.sale_id=sale.id
  if plan_due:db.session.add(CustomerTask(customer_id=customer.id,sale_id=sale.id,task_type='요금제 변경',title=f'{name} 요금제 변경',description=f"{request.form.get('current_plan','')} → {request.form.get('next_plan','')}",due_date=plan_due,assigned_staff=sale.assigned_staff,auto_created=True))
  names=request.form.getlist('addon_name[]');rules=request.form.getlist('addon_rule[]')
  for addon_name,rule in zip(names,rules):
   addon_name=addon_name.strip()
   if not addon_name:continue
   due=due_from_rule(opening,rule);db.session.add(SaleAddon(sale_id=sale.id,name=addon_name,retention_rule=rule,cancellation_due_date=due))
   if due:db.session.add(CustomerTask(customer_id=customer.id,sale_id=sale.id,task_type='부가서비스 해지',title=f'{name} 부가서비스 해지',description=f'{addon_name} · {rule}',due_date=due,assigned_staff=sale.assigned_staff,auto_created=True))
  if internet_due:db.session.add(CustomerTask(customer_id=customer.id,sale_id=sale.id,task_type='인터넷 해지',title=f'{name} 인터넷 해지',description=request.form.get('internet_carrier',''),due_date=internet_due,assigned_staff=sale.assigned_staff,auto_created=True))
  if payback>0:
   db.session.add(Payback(sale_id=sale.id,customer_id=customer.id,amount=payback,due_date=payback_due,status='처리예정',bank=request.form.get('bank'),account_number=request.form.get('account_number'),account_holder=request.form.get('account_holder'),memo=request.form.get('payback_memo')))
   if payback_due:db.session.add(CustomerTask(customer_id=customer.id,sale_id=sale.id,task_type='페이백 지급',title=f'{name} 페이백 지급',description=f'{payback:,}원',due_date=payback_due,assigned_staff=sale.assigned_staff,auto_created=True))
  db.session.commit();flash('개통 등록이 완료되었습니다. 재고·판매일보·고객약속·페이백이 자동 반영되었습니다.','success');return redirect(url_for('sales'))
 return render_template('sale_form.html',staff=staff,partners=partners,branches=branches,today=date.today().isoformat(),sale=None,plans=plans,plan_data=plan_data)


def _customer_for_sale(sale):
 if sale.customer_phone:
  c=Customer.query.filter_by(phone=sale.customer_phone).first()
  if c:return c
 return Customer.query.filter_by(name=sale.customer_name).order_by(Customer.id.desc()).first()

def _rebuild_sale_automation(sale,form):
 customer=_customer_for_sale(sale)
 CustomerTask.query.filter_by(sale_id=sale.id,auto_created=True).delete(synchronize_session=False)
 SaleAddon.query.filter_by(sale_id=sale.id).delete(synchronize_session=False)
 payback=Payback.query.filter_by(sale_id=sale.id).first()
 opening=sale.opening_date
 if sale.plan_change_due_date and sale.next_plan:
  db.session.add(CustomerTask(customer_id=customer.id if customer else None,sale_id=sale.id,task_type='요금제 변경',title=f'{sale.customer_name} 요금제 변경',description=f'{sale.current_plan or ""} → {sale.next_plan or ""}',due_date=sale.plan_change_due_date,assigned_staff=sale.assigned_staff,auto_created=True))
 names=form.getlist('addon_name[]'); rules=form.getlist('addon_rule[]')
 for addon_name,rule in zip(names,rules):
  addon_name=addon_name.strip()
  if not addon_name: continue
  due=due_from_rule(opening,rule); db.session.add(SaleAddon(sale_id=sale.id,name=addon_name,retention_rule=rule,cancellation_due_date=due))
  if due: db.session.add(CustomerTask(customer_id=customer.id if customer else None,sale_id=sale.id,task_type='부가서비스 해지',title=f'{sale.customer_name} 부가서비스 해지',description=f'{addon_name} · {rule}',due_date=due,assigned_staff=sale.assigned_staff,auto_created=True))
 if sale.customer_payback>0:
  if not payback:
   payback=Payback(sale_id=sale.id,customer_id=customer.id if customer else None)
   db.session.add(payback)
  payback.amount=sale.customer_payback; payback.due_date=sale.payback_due_date; payback.bank=form.get('bank'); payback.account_number=form.get('account_number'); payback.account_holder=form.get('account_holder'); payback.memo=form.get('payback_memo')
  if payback.status not in ['완료','취소']: payback.status='처리예정'
  if sale.payback_due_date:
   db.session.add(CustomerTask(customer_id=customer.id if customer else None,sale_id=sale.id,task_type='페이백 지급',title=f'{sale.customer_name} 페이백 지급',description=f'{sale.customer_payback:,}원',due_date=sale.payback_due_date,assigned_staff=sale.assigned_staff,status='완료' if payback.status=='완료' else '처리예정',auto_created=True))
 elif payback:
  db.session.delete(payback)

@app.route('/sales/<int:sid>/edit',methods=['GET','POST'])
@login_required
def sale_edit(sid):
 sale=Sale.query.get_or_404(sid)
 enforce_branch(sale.branch_id)
 staff=(User.query.filter_by(active=True,branch_id=current_branch_id()).order_by(User.display_name).all() if not is_admin() else User.query.filter_by(active=True).order_by(User.display_name,User.username).all()); partners=Partner.query.filter_by(active=True).order_by(Partner.name).all(); branches=(Branch.query.filter_by(id=current_branch_id()).all() if not is_admin() else Branch.query.filter_by(active=True).order_by(Branch.id).all()); plans=PlanMaster.query.filter_by(active=True).order_by(PlanMaster.carrier,PlanMaster.sort_order,PlanMaster.name).all(); plan_data=[{'carrier':p.carrier,'name':p.name} for p in plans]
 if request.method=='POST':
  old_inv=Inventory.query.get(sale.inventory_id) if sale.inventory_id else None
  serial=request.form.get('serial_number','').strip(); new_inv=Inventory.query.filter_by(serial_number=serial).first() if serial else None
  if new_inv and new_inv.sale_id not in [None,sale.id]:
   flash('다른 판매건에 연결된 일련번호입니다.','error'); return redirect(url_for('sale_edit',sid=sale.id))
  if old_inv and (not new_inv or old_inv.id!=new_inv.id):
   old_inv.status='보유중'; old_inv.sale_id=None
   db.session.add(InventoryMovement(inventory_id=old_inv.id,action='판매수정-재고복구',from_branch_id=old_inv.branch_id,to_branch_id=old_inv.branch_id,from_status='판매완료',to_status='보유중',processed_by=session.get('display_name') or session.get('username'),memo=f'판매 #{sale.id} 수정'))
  if new_inv:
   new_inv.status='판매완료'; new_inv.sale_id=sale.id
   if not old_inv or old_inv.id!=new_inv.id:
    db.session.add(InventoryMovement(inventory_id=new_inv.id,action='판매연결',from_branch_id=new_inv.branch_id,to_branch_id=new_inv.branch_id,from_status='보유중',to_status='판매완료',processed_by=session.get('display_name') or session.get('username'),memo=f'판매 #{sale.id} 수정'))
  sale.customer_name=request.form.get('customer_name','').strip() or sale.customer_name; sale.customer_phone=request.form.get('customer_phone','').strip(); sale.customer_birth=request.form.get('customer_birth')
  sale.opening_date=parse_date(request.form.get('opening_date')) or sale.opening_date; sale.carrier=request.form.get('carrier'); sale.opening_type=request.form.get('opening_type'); sale.visit_source=request.form.get('visit_source')
  sale.branch_id=(sale.branch_id if not is_admin() else (request.form.get('branch_id') or sale.branch_id)); sale.assigned_staff=request.form.get('assigned_staff'); sale.serial_number=serial; sale.inventory_id=new_inv.id if new_inv else None; sale.partner_id=(new_inv.partner_id if new_inv else (request.form.get('partner_id') or None))
  sale.manufacturer=new_inv.manufacturer if new_inv else request.form.get('manufacturer'); sale.device=new_inv.model if new_inv else request.form.get('device'); sale.storage=new_inv.capacity if new_inv else request.form.get('storage'); sale.color=new_inv.color if new_inv else request.form.get('color')
  sale.current_plan=request.form.get('current_plan'); sale.next_plan=request.form.get('next_plan'); sale.plan=sale.current_plan; sale.plan_change_due_date=sale.opening_date+timedelta(days=183) if sale.next_plan else None
  sale.internet_carrier=None; sale.internet_subscriber=None; sale.internet_install_date=None; sale.internet_cancel_due_date=None
  sale.payback_due_date=parse_date(request.form.get('payback_due_date')); sale.memo=request.form.get('memo')
  sale.rebate=money(request.form.get('rebate')); sale.verbal_extra=money(request.form.get('verbal_extra')); sale.deduction=money(request.form.get('deduction')); sale.extra_support=money(request.form.get('extra_support')); sale.customer_payback=money(request.form.get('customer_payback')); sale.sim_payment_type=request.form.get('sim_payment_type','없음'); sale.sim_fee=7700
  settlement,tax,margin,transfer_fee=calc_settlement(sale.rebate,sale.verbal_extra,sale.deduction,sale.extra_support,sale.customer_payback,sale.opening_type,sale.sim_payment_type)
  sale.settlement_amount_v2=settlement; sale.tax_amount=tax; sale.final_margin=margin; sale.transfer_fee=transfer_fee; sale.settlement=str(settlement); sale.margin=str(margin)
  customer=_customer_for_sale(sale)
  if customer:
   customer.name=sale.customer_name; customer.phone=sale.customer_phone; customer.carrier=sale.carrier; customer.status='개통고객'
  _rebuild_sale_automation(sale,request.form); db.session.commit()
  flash('판매일보가 수정되었고 고객약속·페이백·재고 연결도 함께 갱신되었습니다.','success'); return redirect(url_for('sales'))
 addons=SaleAddon.query.filter_by(sale_id=sale.id).order_by(SaleAddon.id).all(); pb=Payback.query.filter_by(sale_id=sale.id).first()
 return render_template('sale_edit.html',sale=sale,addons=addons,payback=pb,staff=staff,partners=partners,branches=branches,plans=plans,plan_data=plan_data)

@app.route('/sales')
@login_required
def sales():
 q=request.args.get('q','').strip(); day=request.args.get('date','').strip(); branch_id=request.args.get('branch_id','').strip()
 query=Sale.query
 if not is_admin():
  branch_id=str(current_branch_id() or ''); query=query.filter(Sale.branch_id==current_branch_id()) if current_branch_id() else query.filter(Sale.id==-1)
 elif branch_id:
  try:query=query.filter(Sale.branch_id==int(branch_id))
  except:pass
 if day:
  d=parse_date(day)
  if d: query=query.filter(Sale.opening_date==d)
 if q:query=query.filter(or_(Sale.customer_name.ilike(f'%{q}%'),Sale.customer_phone.ilike(f'%{q}%'),Sale.device.ilike(f'%{q}%'),Sale.serial_number.ilike(f'%{q}%')))
 items=query.order_by(Sale.opening_date.desc(),Sale.id.desc()).all()
 doc_counts=dict(db.session.query(SaleDocument.sale_id,db.func.count(SaleDocument.id)).filter(SaleDocument.sale_id.in_([s.id for s in items] or [0])).group_by(SaleDocument.sale_id).all())
 return render_template('sales.html',sales=items,q=q,date_filter=day,branch_id=branch_id,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all(),doc_counts=doc_counts,total_settlement=sum(s.settlement_amount_v2 or money(s.settlement) for s in items),total_margin=sum(s.final_margin or money(s.margin) for s in items))


@app.route('/sales/<int:sid>/documents',methods=['GET','POST'])
@login_required
def sale_documents(sid):
 sale=Sale.query.get_or_404(sid); enforce_branch(sale.branch_id)
 if request.method=='POST':
  f=request.files.get('file')
  if not f or not f.filename:
   flash('저장할 서류 파일을 선택해주세요.','error'); return redirect(url_for('sale_documents',sid=sid))
  name=secure_filename(f.filename) or 'document'
  content_type=(f.mimetype or '').lower()
  allowed={'application/pdf','image/jpeg','image/png','image/webp'}
  if content_type not in allowed:
   flash('PDF, JPG, PNG, WEBP 파일만 저장할 수 있습니다.','error'); return redirect(url_for('sale_documents',sid=sid))
  data=f.read()
  if len(data)>10*1024*1024:
   flash('서류 1개는 10MB 이하만 저장할 수 있습니다.','error'); return redirect(url_for('sale_documents',sid=sid))
  db.session.add(SaleDocument(sale_id=sale.id,branch_id=sale.branch_id,doc_type=request.form.get('doc_type','기타서류'),original_name=name,content_type=content_type,file_size=len(data),file_data=data,uploaded_by=session.get('display_name') or session.get('username')))
  db.session.commit(); flash('고객 서류가 안전하게 저장되었습니다.','success'); return redirect(url_for('sale_documents',sid=sid))
 docs=SaleDocument.query.filter_by(sale_id=sale.id).order_by(SaleDocument.created_at.desc()).all()
 return render_template('sale_documents.html',sale=sale,docs=docs)

@app.get('/documents/<int:did>/view')
@login_required
def document_view(did):
 d=SaleDocument.query.get_or_404(did); sale=Sale.query.get_or_404(d.sale_id); enforce_branch(sale.branch_id)
 audit('고객서류 열람','sale_document',d.id,f'{sale.customer_name} · {d.doc_type} · {d.original_name}',sale.branch_id);db.session.commit()
 return send_file(io.BytesIO(d.file_data),mimetype=d.content_type,download_name=d.original_name,as_attachment=False)

@app.post('/documents/<int:did>/delete')
@login_required
def document_delete(did):
 d=SaleDocument.query.get_or_404(did); sale=Sale.query.get_or_404(d.sale_id); enforce_branch(sale.branch_id)
 db.session.delete(d); db.session.commit(); flash('서류가 삭제되었습니다.','success'); return redirect(url_for('sale_documents',sid=sale.id))

@app.post('/sales/<int:sid>/delete')
@login_required
@admin_required
def sale_delete(sid):
 s=Sale.query.get_or_404(sid)
 inv=Inventory.query.filter_by(sale_id=s.id).first()
 if inv:
  old_status=inv.status; inv.status='보유중'; inv.sale_id=None
  db.session.add(InventoryMovement(inventory_id=inv.id,action='판매삭제-재고복구',from_branch_id=inv.branch_id,to_branch_id=inv.branch_id,from_status=old_status,to_status='보유중',processed_by=session.get('display_name') or session.get('username'),memo=f'판매 #{s.id} 삭제'))
 CustomerTask.query.filter_by(sale_id=s.id).delete(synchronize_session=False)
 Payback.query.filter_by(sale_id=s.id).delete(synchronize_session=False)
 SaleAddon.query.filter_by(sale_id=s.id).delete(synchronize_session=False)
 SaleDocument.query.filter_by(sale_id=s.id).delete(synchronize_session=False)
 db.session.delete(s); db.session.commit(); flash('개통건이 삭제되었고 연결 재고는 보유중으로 복구되었습니다.','success'); return redirect(url_for('sales'))

@app.route('/paybacks')
@login_required
def paybacks():
 status=request.args.get('status','').strip(); due=request.args.get('due','').strip(); branch_id=request.args.get('branch_id','').strip()
 query=Payback.query.join(Sale,Payback.sale_id==Sale.id)
 if not is_admin():
  branch_id=str(current_branch_id() or ''); query=query.filter(Sale.branch_id==current_branch_id()) if current_branch_id() else query.filter(Payback.id==-1)
 elif branch_id:
  try:query=query.filter(Sale.branch_id==int(branch_id))
  except:pass
 if status: query=query.filter(Payback.status==status)
 today=date.today()
 if due=='today': query=query.filter(Payback.due_date==today,Payback.status!='완료')
 elif due=='overdue': query=query.filter(Payback.due_date<today,Payback.status!='완료')
 items=query.order_by(Payback.due_date.asc(),Payback.id.desc()).all()
 sales_map={s.id:s for s in Sale.query.filter(Sale.id.in_([p.sale_id for p in items] or [0])).all()}
 allq=payback_query_scoped()
 stats={'today':allq.filter(Payback.due_date==today,Payback.status!='완료').count(),'overdue':payback_query_scoped().filter(Payback.due_date<today,Payback.status!='완료').count(),'pending':payback_query_scoped().filter(Payback.status!='완료').count(),'approval':payback_query_scoped().filter(Payback.approval_status=='승인대기',Payback.status!='완료').count()}
 return render_template('paybacks.html',items=items,sales_map=sales_map,status=status,due=due,branch_id=branch_id,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all(),stats=stats)

@app.route('/paybacks/<int:pid>/edit',methods=['GET','POST'])
@login_required
def payback_edit(pid):
 p=Payback.query.get_or_404(pid); s=Sale.query.get(p.sale_id)
 if s: enforce_branch(s.branch_id)
 if request.method=='POST':
  old_sensitive=(p.amount,p.bank,p.account_number,p.account_holder)
  p.amount=money(request.form.get('amount')); p.due_date=parse_date(request.form.get('due_date')); p.status=request.form.get('status','처리예정')
  p.bank=request.form.get('bank'); p.account_number=request.form.get('account_number'); p.account_holder=request.form.get('account_holder'); p.memo=request.form.get('memo')
  if old_sensitive!=(p.amount,p.bank,p.account_number,p.account_holder):p.approval_status='승인대기';p.approved_at=None;p.approved_by=None;p.rejection_reason=None
  if p.status=='완료':
   if p.approval_status!='승인':flash('승인되지 않은 페이백은 지급완료 처리할 수 없습니다.','error');return redirect(url_for('payback_edit',pid=pid))
   if not p.processed_at: p.processed_at=datetime.utcnow()
   p.processed_by=session.get('display_name') or session.get('username')
  else:
   p.processed_at=None; p.processed_by=None
  if s:
   s.customer_payback=p.amount; s.payback_due_date=p.due_date
   settlement,tax,margin,transfer_fee=calc_settlement(s.rebate or 0,s.verbal_extra or 0,s.deduction or 0,s.extra_support or 0,p.amount,s.opening_type,s.sim_payment_type)
   s.settlement_amount_v2=settlement; s.tax_amount=tax; s.final_margin=margin; s.transfer_fee=transfer_fee; s.settlement=str(settlement); s.margin=str(margin)
  task=CustomerTask.query.filter_by(sale_id=p.sale_id,task_type='페이백 지급',auto_created=True).order_by(CustomerTask.id.desc()).first()
  if p.amount>0 and p.due_date:
   if not task:
    customer=_customer_for_sale(s) if s else None
    task=CustomerTask(customer_id=customer.id if customer else None,sale_id=p.sale_id,task_type='페이백 지급',title=f'{s.customer_name if s else "고객"} 페이백 지급',description=f'{p.amount:,}원',due_date=p.due_date,assigned_staff=s.assigned_staff if s else None,auto_created=True)
    db.session.add(task)
   task.due_date=p.due_date; task.description=f'{p.amount:,}원'; task.status='완료' if p.status=='완료' else '처리예정'
   task.completed_at=p.processed_at if p.status=='완료' else None; task.completed_by=p.processed_by if p.status=='완료' else None
  elif task:
   db.session.delete(task)
  db.session.commit(); flash('페이백 정보와 연결된 고객약속/최종마진이 수정되었습니다.','success'); return redirect(url_for('paybacks'))
 return render_template('payback_edit.html',p=p,sale=s)

@app.post('/paybacks/<int:pid>/complete')
@login_required
def payback_complete(pid):
 p=Payback.query.get_or_404(pid); s=Sale.query.get_or_404(p.sale_id); enforce_branch(s.branch_id)
 if p.approval_status!='승인':flash('지정 승인권자의 승인이 필요합니다.','error');return redirect(url_for('paybacks'))
 p.status='완료'; p.processed_at=datetime.utcnow(); p.processed_by=session.get('display_name') or session.get('username')
 task=CustomerTask.query.filter_by(sale_id=p.sale_id,task_type='페이백 지급',auto_created=True).order_by(CustomerTask.id.desc()).first()
 if task: task.status='완료'; task.completed_at=p.processed_at; task.completed_by=p.processed_by
 audit('페이백 지급완료','payback',p.id,f'{p.amount}원 · {p.bank} · {p.account_number}',s.branch_id);db.session.commit(); flash('페이백을 완료 처리했습니다.','success'); return redirect(url_for('paybacks'))

@app.post('/paybacks/<int:pid>/approval')
@login_required
def payback_approval(pid):
 if not can_approve_payback():abort(403)
 p=Payback.query.get_or_404(pid);s=Sale.query.get_or_404(p.sale_id);enforce_branch(s.branch_id);decision=request.form.get('decision')
 if decision=='approve':p.approval_status='승인';p.approved_at=datetime.utcnow();p.approved_by=session.get('display_name') or session.get('username');p.rejection_reason=None;msg='페이백 지급을 승인했습니다.'
 elif decision=='reject':p.approval_status='반려';p.approved_at=None;p.approved_by=session.get('display_name') or session.get('username');p.rejection_reason=request.form.get('reason','').strip() or '정보 재확인 필요';msg='페이백 지급을 반려했습니다.'
 else:abort(400)
 audit('페이백 '+p.approval_status,'payback',p.id,f'{p.amount}원 · {p.rejection_reason or ""}',s.branch_id);db.session.commit();flash(msg,'success');return redirect(request.referrer or url_for('paybacks'))

@app.post('/paybacks/bulk-transfer.xlsx')
@login_required
def payback_bulk_transfer():
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 ids=[]
 for v in request.form.getlist('payback_ids'):
  try:ids.append(int(v))
  except:pass
 items=payback_query_scoped().filter(Payback.id.in_(ids or [0]),Payback.approval_status=='승인',Payback.status!='완료').order_by(Payback.id).all()
 if not items:flash('승인된 미지급 페이백을 선택해주세요.','error');return redirect(url_for('paybacks'))
 sales={s.id:s for s in Sale.query.filter(Sale.id.in_([p.sale_id for p in items])).all()};month=date.today().strftime('%m월');wb=Workbook();ws=wb.active;ws.title='공통 대량이체';headers=['고객명','은행명','계좌번호','금액','받는통장 표시','보내는통장 표시','검증상태']
 ws.append(headers)
 for p in items:
  s=sales.get(p.sale_id);valid='정상' if all([p.bank,p.account_number,p.account_holder,p.amount>0]) else '확인필요';ws.append([s.customer_name if s else p.account_holder,p.bank or '',p.account_number or '',p.amount,'고무신모바일',f'{month} {s.customer_name if s else p.account_holder}',valid])
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='14324A');cell.alignment=Alignment(horizontal='center')
 ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:G{ws.max_row}'
 for col,w in zip('ABCDEFG',[16,13,24,15,20,22,13]):ws.column_dimensions[col].width=w
 for title in ['우리은행 업로드','KB국민은행 업로드']:
  bank_ws=wb.create_sheet(title);bank_ws.append(['은행코드/은행명','계좌번호','이체금액','받는분 통장표시','내 통장표시'])
  for p in items:
   s=sales.get(p.sale_id);bank_ws.append([p.bank or '',p.account_number or '',p.amount,'고무신모바일',f'{month} {s.customer_name if s else p.account_holder}'])
  bank_ws.freeze_panes='A2'
 audit('페이백 대량이체 다운로드','payback','bulk',f'{len(items)}건 / {sum(p.amount for p in items):,}원');db.session.commit();out=io.BytesIO();wb.save(out);out.seek(0)
 return send_file(out,as_attachment=True,download_name=f'TrustFlow_페이백대량이체_{date.today()}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.post('/paybacks/<int:pid>/reopen')
@login_required
def payback_reopen(pid):
 p=Payback.query.get_or_404(pid); enforce_branch(Sale.query.get_or_404(p.sale_id).branch_id); p.status='처리예정'; p.processed_at=None; p.processed_by=None
 task=CustomerTask.query.filter_by(sale_id=p.sale_id,task_type='페이백 지급',auto_created=True).order_by(CustomerTask.id.desc()).first()
 if task: task.status='처리예정'; task.completed_at=None; task.completed_by=None
 db.session.commit(); flash('페이백을 처리예정으로 되돌렸습니다.','success'); return redirect(url_for('paybacks'))


@app.route('/wired-sales')
@login_required
def wired_sales():
 q=request.args.get('q','').strip(); status=request.args.get('status','').strip(); branch_id=request.args.get('branch_id','').strip()
 query=WiredSale.query
 if not is_admin():
  branch_id=str(current_branch_id() or ''); query=query.filter(WiredSale.branch_id==current_branch_id()) if current_branch_id() else query.filter(WiredSale.id==-1)
 elif branch_id:
  try: query=query.filter(WiredSale.branch_id==int(branch_id))
  except: pass
 if q: query=query.filter(or_(WiredSale.customer_name.ilike(f'%{q}%'),WiredSale.customer_phone.ilike(f'%{q}%'),WiredSale.subscriber_name.ilike(f'%{q}%')))
 if status: query=query.filter_by(status=status)
 items=query.order_by(WiredSale.sale_date.desc(),WiredSale.id.desc()).all()
 branches={b.id:b for b in Branch.query.all()}
 return render_template('wired_sales.html',items=items,branches=branches,branch_list=Branch.query.filter_by(active=True).order_by(Branch.id).all(),q=q,status=status,branch_id=branch_id)

@app.route('/wired-sales/new',methods=['GET','POST'])
@login_required
def wired_sale_new():
 branches=(Branch.query.filter_by(id=current_branch_id()).all() if not is_admin() else Branch.query.filter_by(active=True).order_by(Branch.id).all())
 staff=(User.query.filter_by(active=True,branch_id=current_branch_id()).order_by(User.display_name).all() if not is_admin() else User.query.filter_by(active=True).order_by(User.display_name,User.username).all())
 wired_products=WiredProductMaster.query.filter_by(active=True).order_by(WiredProductMaster.carrier,WiredProductMaster.category,WiredProductMaster.sort_order,WiredProductMaster.name).all()
 wired_product_data=[{'carrier':x.carrier,'category':x.category,'name':x.name,'speed':x.speed or ''} for x in wired_products]
 if request.method=='POST':
  rebate=money(request.form.get('rebate')); verbal=money(request.form.get('verbal_extra'))
  gift=money(request.form.get('gift_certificate')); gift_cost=money(request.form.get('gift_cost')); deduction=money(request.form.get('deduction'))
  payback=money(request.form.get('payback')); gift_return=money(request.form.get('gift_return'))
  settlement,tax,margin=calc_wired_settlement(rebate,verbal,gift,gift_cost,deduction,payback,gift_return)
  item=WiredSale(
   sale_date=parse_date(request.form.get('sale_date')) or date.today(),
   customer_name=request.form.get('customer_name','').strip(),
   customer_phone=request.form.get('customer_phone','').strip(),
   subscriber_name=request.form.get('subscriber_name','').strip(),
   branch_id=(current_branch_id() if not is_admin() else (request.form.get('branch_id') or None)),
   assigned_staff=request.form.get('assigned_staff') or session.get('display_name') or session.get('username'),
   carrier=request.form.get('carrier'),
   business_type=request.form.get('business_type','유선판매'),
   product_type=request.form.get('product_type'),
   internet_plan=request.form.get('internet_plan'),
   internet_speed=request.form.get('internet_speed'),
   tv_plan=request.form.get('tv_plan'),
   install_due_date=parse_date(request.form.get('install_due_date')),
   install_date=parse_date(request.form.get('install_date')),
   status=request.form.get('status','접수'),
   rebate=rebate,verbal_extra=verbal,gift_certificate=gift,gift_cost=gift_cost,deduction=deduction,
   settlement_amount=settlement,tax_rate=.133,tax_amount=tax,
   payback=payback,gift_return=gift_return,final_margin=margin,
   memo=request.form.get('memo'),
   created_by=session.get('display_name') or session.get('username')
  )
  if not item.customer_name:
   flash('고객명을 입력해주세요.','error')
   return render_template('wired_sale_form.html',item=None,branches=branches,staff=staff,today=date.today().isoformat(),wired_product_data=wired_product_data)
  db.session.add(item); db.session.commit()
  flash('유선판매 내역이 등록되었습니다.','success'); return redirect(url_for('wired_sales'))
 return render_template('wired_sale_form.html',item=None,branches=branches,staff=staff,today=date.today().isoformat(),wired_product_data=wired_product_data)

@app.route('/wired-sales/<int:wid>/edit',methods=['GET','POST'])
@login_required
def wired_sale_edit(wid):
 item=WiredSale.query.get_or_404(wid)
 enforce_branch(item.branch_id)
 branches=(Branch.query.filter_by(id=current_branch_id()).all() if not is_admin() else Branch.query.filter_by(active=True).order_by(Branch.id).all())
 staff=(User.query.filter_by(active=True,branch_id=current_branch_id()).order_by(User.display_name).all() if not is_admin() else User.query.filter_by(active=True).order_by(User.display_name,User.username).all())
 wired_products=WiredProductMaster.query.filter_by(active=True).order_by(WiredProductMaster.carrier,WiredProductMaster.category,WiredProductMaster.sort_order,WiredProductMaster.name).all()
 wired_product_data=[{'carrier':x.carrier,'category':x.category,'name':x.name,'speed':x.speed or ''} for x in wired_products]
 if request.method=='POST':
  item.sale_date=parse_date(request.form.get('sale_date')) or item.sale_date
  item.customer_name=request.form.get('customer_name','').strip() or item.customer_name
  item.customer_phone=request.form.get('customer_phone','').strip()
  item.subscriber_name=request.form.get('subscriber_name','').strip()
  item.branch_id=request.form.get('branch_id') or None
  item.assigned_staff=request.form.get('assigned_staff')
  item.carrier=request.form.get('carrier')
  item.business_type=request.form.get('business_type','유선판매')
  item.product_type=request.form.get('product_type')
  item.internet_plan=request.form.get('internet_plan')
  item.internet_speed=request.form.get('internet_speed')
  item.tv_plan=request.form.get('tv_plan')
  item.install_due_date=parse_date(request.form.get('install_due_date'))
  item.install_date=parse_date(request.form.get('install_date'))
  item.status=request.form.get('status','접수')
  item.rebate=money(request.form.get('rebate')); item.verbal_extra=money(request.form.get('verbal_extra'))
  item.gift_certificate=money(request.form.get('gift_certificate')); item.gift_cost=money(request.form.get('gift_cost')); item.deduction=money(request.form.get('deduction'))
  item.payback=money(request.form.get('payback')); item.gift_return=money(request.form.get('gift_return'))
  item.settlement_amount,item.tax_amount,item.final_margin=calc_wired_settlement(
   item.rebate,item.verbal_extra,item.gift_certificate,item.gift_cost,item.deduction,item.payback,item.gift_return
  )
  item.tax_rate=.133; item.memo=request.form.get('memo')
  db.session.commit()
  flash('유선판매 내역과 정산정보가 수정되었습니다.','success'); return redirect(url_for('wired_sales'))
 return render_template('wired_sale_form.html',item=item,branches=branches,staff=staff,today=date.today().isoformat(),wired_product_data=wired_product_data)

@app.post('/wired-sales/<int:wid>/delete')
@login_required
@admin_required
def wired_sale_delete(wid):
 item=WiredSale.query.get_or_404(wid); enforce_branch(item.branch_id); db.session.delete(item); db.session.commit(); flash('유선판매 내역이 삭제되었습니다.','success'); return redirect(url_for('wired_sales'))


@app.route('/masters',methods=['GET','POST'])
@login_required
@admin_required
def masters():
 prepare_database()
 if request.method=='POST':
  kind=request.form.get('kind')
  if kind=='plan':
   carrier=request.form.get('carrier','').strip(); name=request.form.get('name','').strip()
   if carrier and name and not PlanMaster.query.filter_by(carrier=carrier,name=name).first():
    db.session.add(PlanMaster(carrier=carrier,name=name,active=True,sort_order=999)); db.session.commit(); flash('요금제가 추가되었습니다.','success')
  elif kind=='device':
   maker=request.form.get('manufacturer','').strip(); model=request.form.get('model','').strip()
   if maker and model and not DeviceMaster.query.filter_by(manufacturer=maker,model=model).first():
    db.session.add(DeviceMaster(manufacturer=maker,model=model,capacities=request.form.get('capacities'),colors=request.form.get('colors'),active=True,sort_order=999)); db.session.commit(); flash('단말기 기준정보가 추가되었습니다.','success')
  elif kind=='wired':
   carrier=request.form.get('carrier','').strip(); category=request.form.get('category','').strip(); name=request.form.get('name','').strip(); speed=request.form.get('speed','').strip() or None
   if carrier and category and name and not WiredProductMaster.query.filter_by(carrier=carrier,category=category,name=name).first():
    db.session.add(WiredProductMaster(carrier=carrier,category=category,name=name,speed=speed,active=True,sort_order=999)); db.session.commit(); flash('유선상품 기준정보가 추가되었습니다.','success')
  return redirect(url_for('masters'))
 return render_template('masters.html',plans=PlanMaster.query.order_by(PlanMaster.carrier,PlanMaster.sort_order,PlanMaster.name).all(),devices=DeviceMaster.query.order_by(DeviceMaster.manufacturer,DeviceMaster.sort_order,DeviceMaster.model).all(),wired_products=WiredProductMaster.query.order_by(WiredProductMaster.carrier,WiredProductMaster.category,WiredProductMaster.sort_order,WiredProductMaster.name).all())

@app.post('/masters/plan/<int:mid>/toggle')
@login_required
@admin_required
def master_plan_toggle(mid):
 x=PlanMaster.query.get_or_404(mid); x.active=not x.active; db.session.commit(); return redirect(url_for('masters'))

@app.post('/masters/device/<int:mid>/toggle')
@login_required
@admin_required
def master_device_toggle(mid):
 x=DeviceMaster.query.get_or_404(mid); x.active=not x.active; db.session.commit(); return redirect(url_for('masters'))

@app.post('/masters/wired/<int:mid>/toggle')
@login_required
@admin_required
def master_wired_toggle(mid):
 x=WiredProductMaster.query.get_or_404(mid); x.active=not x.active; db.session.commit(); return redirect(url_for('masters'))

@app.route('/staff',methods=['GET','POST'])
@login_required
@admin_required
def staff():
 prepare_database()
 if request.method=='POST':
  u=request.form.get('username','').strip(); pw=request.form.get('password',''); role=request.form.get('role','staff'); display_name=request.form.get('display_name','').strip(); phone=normalize_phone(request.form.get('recovery_phone','')); branch_id=request.form.get('branch_id') or None; company=session.get('company_code') or 'trustflow'
  if not u or not pw or not display_name:
   flash('직원명, 로그인 아이디, 비밀번호를 모두 입력해주세요.','error')
  elif User.query.filter_by(username=u).first():
   flash('이미 사용 중인 로그인 아이디입니다.','error')
  else:
   db.session.add(User(username=u,password_hash=generate_password_hash(pw),role=role,display_name=display_name,branch_id=branch_id,company_code=company,recovery_phone=phone,active=True)); db.session.commit(); flash('직원이 등록되었습니다.','success')
  return redirect(url_for('staff'))
 company=session.get('company_code') or 'trustflow'
 users=User.query.filter_by(company_code=company).order_by(User.active.desc(),User.display_name,User.username).all()
 account_requests=AccountRequest.query.filter_by(company_code=company,status='대기').order_by(AccountRequest.id.desc()).all()
 return render_template('staff.html',users=users,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all(),account_requests=account_requests)

@app.route('/staff/<int:uid>/edit',methods=['GET','POST'])
@login_required
@admin_required
def staff_edit(uid):
 u=User.query.get_or_404(uid); enforce_user_company(u)
 if request.method=='POST':
  display_name=request.form.get('display_name','').strip()
  if not display_name:
   flash('직원명을 입력해주세요.','error'); return redirect(url_for('staff_edit',uid=uid))
  u.display_name=display_name; u.branch_id=request.form.get('branch_id') or None; u.role=request.form.get('role','staff'); u.recovery_phone=normalize_phone(request.form.get('recovery_phone','')); u.active=request.form.get('active')=='1';u.can_approve_payback=request.form.get('can_approve_payback')=='1'
  new_pw=request.form.get('password','')
  if new_pw: u.password_hash=generate_password_hash(new_pw)
  db.session.commit()
  if session.get('user_id')==u.id:
   session['display_name']=u.display_name; session['role']=u.role
  flash('직원정보가 수정되었습니다.','success'); return redirect(url_for('staff'))
 return render_template('staff_edit.html',u=u,branches=Branch.query.filter_by(active=True).order_by(Branch.id).all())

@app.post('/staff/<int:uid>/toggle')
@login_required
@admin_required
def staff_toggle(uid):
 u=User.query.get_or_404(uid); enforce_user_company(u)
 if session.get('user_id')==u.id and u.active:
  flash('현재 로그인 중인 본인 계정은 비활성화할 수 없습니다.','error'); return redirect(url_for('staff'))
 u.active=not bool(u.active); db.session.commit()
 flash('직원 계정 상태를 변경했습니다.','success'); return redirect(url_for('staff'))

@app.post('/staff/<int:uid>/delete')
@login_required
@admin_required
def staff_delete(uid):
 u=User.query.get_or_404(uid); enforce_user_company(u)
 if session.get('user_id')==u.id:
  flash('현재 로그인 중인 본인 계정은 삭제할 수 없습니다.','error'); return redirect(url_for('staff'))
 staff_name=u.display_name or u.username
 has_sales=Sale.query.filter(or_(Sale.assigned_staff==staff_name,Sale.created_by==staff_name,Sale.assigned_staff==u.username,Sale.created_by==u.username)).count()>0
 has_tasks=CustomerTask.query.filter(or_(CustomerTask.assigned_staff==staff_name,CustomerTask.completed_by==staff_name,CustomerTask.assigned_staff==u.username,CustomerTask.completed_by==u.username)).count()>0
 has_wired=WiredSale.query.filter(or_(WiredSale.assigned_staff==staff_name,WiredSale.created_by==staff_name,WiredSale.assigned_staff==u.username,WiredSale.created_by==u.username)).count()>0
 if has_sales or has_tasks or has_wired:
  flash('판매/약속/유선판매 이력이 있는 직원은 완전 삭제할 수 없습니다. 비활성화로 관리해주세요.','error'); return redirect(url_for('staff'))
 db.session.delete(u); db.session.commit(); flash('직원 계정이 삭제되었습니다.','success'); return redirect(url_for('staff'))

@app.post('/account-requests/<int:request_id>/complete')
@login_required
@admin_required
def account_request_complete(request_id):
 item=AccountRequest.query.get_or_404(request_id)
 if item.company_code!=(session.get('company_code') or 'trustflow'): abort(403)
 item.status='완료'; db.session.commit(); flash('비밀번호 재설정 요청을 완료 처리했습니다.','success'); return redirect(url_for('staff'))

@app.route('/audit-logs')
@login_required
@admin_required
def audit_logs():
 prepare_database();action=request.args.get('action','').strip();day=request.args.get('date','').strip();q=AuditLog.query.filter_by(company_code=session.get('company_code') or 'trustflow')
 if action:q=q.filter_by(action=action)
 if day:
  d=parse_date(day)
  if d:q=q.filter(AuditLog.created_at>=datetime.combine(d,datetime.min.time()),AuditLog.created_at<datetime.combine(d+timedelta(days=1),datetime.min.time()))
 items=q.order_by(AuditLog.created_at.desc()).limit(1000).all();actions=[x[0] for x in db.session.query(AuditLog.action).filter_by(company_code=session.get('company_code') or 'trustflow').distinct().order_by(AuditLog.action).all()]
 return render_template('audit_logs.html',items=items,actions=actions,action=action,date_filter=day,branches={b.id:b for b in Branch.query.all()})

@app.get('/admin/backup.xlsx')
@login_required
@admin_required
def admin_backup():
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 wb=Workbook();wb.remove(wb.active);branch_names={b.id:b.name for b in Branch.query.all()}
 def sheet(title,headers,rows):
  ws=wb.create_sheet(title);ws.append(headers)
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='14324A')
  for row in rows:ws.append(row)
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for col in ws.columns:
   letter=col[0].column_letter;ws.column_dimensions[letter].width=min(38,max(11,max(len(str(c.value or '')) for c in col)+2))
 sheet('고객',['ID','고객명','휴대전화','통신사','상태','도로명주소','지번주소','상세주소','메모','등록일'],[(x.id,x.name,x.phone,x.carrier,x.status,x.address_road,x.address_jibun,x.address_detail,x.memo,x.created_at) for x in Customer.query.order_by(Customer.id).all()])
 sheet('판매일보',['ID','개통일','지점','고객명','휴대전화','통신사','개통유형','단말기','용량','색상','일련번호','요금제','판매자','정산금','최종마진'],[(x.id,x.opening_date,branch_names.get(x.branch_id),x.customer_name,x.customer_phone,x.carrier,x.opening_type,x.device,x.storage,x.color,x.serial_number,x.current_plan,x.assigned_staff,x.settlement_amount_v2,x.final_margin) for x in Sale.query.order_by(Sale.id).all()])
 sheet('재고',['ID','지점','일련번호','통신사','제조사','모델','용량','색상','입고일','상태','매입가'],[(x.id,branch_names.get(x.branch_id),x.serial_number,x.carrier,x.manufacturer,x.model,x.capacity,x.color,x.received_date,x.status,x.purchase_price) for x in Inventory.query.order_by(Inventory.id).all()])
 sheet('페이백',['ID','판매ID','예정일','금액','은행','계좌번호','예금주','승인상태','승인자','지급상태','지급자'],[(x.id,x.sale_id,x.due_date,x.amount,x.bank,x.account_number,x.account_holder,x.approval_status,x.approved_by,x.status,x.processed_by) for x in Payback.query.order_by(Payback.id).all()])
 sheet('시재',['ID','날짜','지점','입출금','항목','수단','금액','거래처','메모','등록자'],[(x.id,x.ledger_date,branch_names.get(x.branch_id),x.direction,x.category,x.payment_method,x.amount,x.counterparty,x.memo,x.created_by) for x in CashLedger.query.order_by(CashLedger.id).all()])
 sheet('상담기록',['ID','고객ID','지점','일시','담당자','채널','결과','내용','다음연락일'],[(x.id,x.customer_id,branch_names.get(x.branch_id),x.contacted_at,x.staff_name,x.channel,x.outcome,x.note,x.next_contact_date) for x in ContactLog.query.order_by(ContactLog.id).all()])
 sheet('법률업무',['ID','고객ID','지점','유형','금액','발생일','청구사유','증빙','지급기한','상태','담당자'],[(x.id,x.customer_id,branch_names.get(x.branch_id),x.case_type,x.claim_amount,x.incident_date,x.reason,x.evidence,x.demand_due_date,x.status,x.assigned_staff) for x in LegalCase.query.order_by(LegalCase.id).all()])
 audit('관리자 전체백업 다운로드','system','backup',f'{date.today()} 운영데이터 7개 시트');db.session.commit();out=io.BytesIO();wb.save(out);out.seek(0)
 return send_file(out,as_attachment=True,download_name=f'TrustFlow_운영백업_{date.today()}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
